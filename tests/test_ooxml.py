from __future__ import annotations

from io import BytesIO
from pathlib import Path
import shutil
import subprocess
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.part import PartFactory
from docx.oxml.ns import qn
from docx.oxml.parser import parse_xml
from docx.oxml.ns import nsdecls
from docx.parts.image import ImagePart
from docx.shared import Emu, Inches, Mm, Pt, Twips
from PIL import Image

from mdtoword import ooxml
from mdtoword.ooxml import (
    FootnoteManager,
    FootnotesPart,
    ListNumbering,
    add_bookmark,
    add_external_hyperlink,
    add_field,
    add_hyperlink_run,
    add_internal_hyperlink,
    add_page_number_footer,
    add_svg_picture,
    add_toc,
    append_toc,
    apply_page_setup,
    bookmark_name,
    detect_language,
    github_slug,
    insert_toc_before,
    normalize_raster,
    picture_size_to_fit,
    request_field_update_on_open,
    set_cant_split,
    set_document_language,
    set_keep_lines,
    set_paragraph_borders,
    set_paragraph_shading,
    set_picture_description,
    set_run_shading,
    set_table_header_repeat,
    svg_intrinsic_size,
    text_height,
    text_width,
)


_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_ASVG = "{http://schemas.microsoft.com/office/drawing/2016/SVG/main}"
_OFFICECLI = shutil.which("officecli")

_SVG = (
    b'<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100" viewBox="0 0 200 100">'
    b'<rect width="200" height="100" fill="#3a7"/></svg>'
)


# What Adobe Illustrator exports: comment, DOCTYPE with an internal subset and
# the namespace itself written as an entity reference.
_ILLUSTRATOR_SVG = (
    b'<?xml version="1.0" encoding="utf-8"?>\n<!-- Generator: Adobe Illustrator 27.0 -->\n'
    b'<!DOCTYPE svg PUBLIC "-//W3C//DTD SVG 1.1//EN" "http://www.w3.org/Graphics/SVG/1.1/DTD/svg11.dtd" [\n'
    b'  <!ENTITY ns_svg "http://www.w3.org/2000/svg">\n  <!ENTITY st0 "fill:#3a7">\n]>\n'
    b'<svg version="1.1" xmlns="&ns_svg;" width="120px" height="60px" viewBox="0 0 120 60">'
    b'<rect width="120" height="60" style="&st0;"/><text>a &amp; b <![CDATA[<x>]]></text></svg>'
)
_EXTERNAL_ENTITY_SVG = (
    b'<?xml version="1.0"?><!DOCTYPE svg [<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
    b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="5"><text>&xxe;</text></svg>'
)
_BILLION_LAUGHS_SVG = (
    b'<?xml version="1.0"?><!DOCTYPE svg [<!ENTITY a "aaaaaaaaaa">'
    + b"".join(
        b'<!ENTITY %s "%s">' % (bytes([98 + i]), b"&%s;" % bytes([97 + i]) * 10) for i in range(8)
    )
    + b']><svg xmlns="http://www.w3.org/2000/svg" width="10" height="5"><text>&i;</text></svg>'
)


def _image_bytes(fmt: str = "PNG", size: tuple[int, int] = (4, 2), mode: str = "RGB") -> bytes:
    out = BytesIO()
    Image.new(mode, size, (200, 30, 30) if mode == "RGB" else (200, 30, 30, 128)).save(out, fmt)
    return out.getvalue()


def _reopen(document):
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)
    return Document(stream), stream.getvalue()


def _zip_text(blob: bytes, name: str) -> str:
    with zipfile.ZipFile(BytesIO(blob)) as archive:
        return archive.read(name).decode("utf-8")


class _OrderAssertions(unittest.TestCase):
    def assert_schema_order(self, parent, sequence) -> None:
        positions = [sequence.index(child.tag) for child in parent if child.tag in sequence]
        self.assertEqual(positions, sorted(positions), [child.tag for child in parent])


class ListNumberingTests(_OrderAssertions):
    def setUp(self) -> None:
        self.document = Document()
        self.numbering = self.document.part.numbering_part.element
        self.lists = ListNumbering(self.document)

    def _num(self, num_id: int):
        return self.numbering.xpath(f'./w:num[@w:numId="{num_id}"]')[0]

    def _abstract_for(self, num_id: int):
        abstract_id = self._num(num_id).find(qn("w:abstractNumId")).get(qn("w:val"))
        return self.numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]')[0]

    def test_ids_are_fresh_and_above_existing_maxima(self) -> None:
        existing_nums = [int(v) for v in self.numbering.xpath("./w:num/@w:numId")]
        existing_abstracts = [int(v) for v in self.numbering.xpath("./w:abstractNum/@w:abstractNumId")]
        first = self.lists.start_list(ordered=True)
        second = self.lists.start_list(ordered=True)
        self.assertEqual(first, max(existing_nums) + 1)
        self.assertEqual(second, first + 1)
        abstract_id = int(self._num(first).find(qn("w:abstractNumId")).get(qn("w:val")))
        self.assertGreater(abstract_id, max(existing_abstracts))

    def test_abstract_definitions_are_shared_and_precede_all_nums(self) -> None:
        before = len(self.numbering.findall(qn("w:abstractNum")))
        ordered_a = self.lists.start_list(ordered=True)
        bullet = self.lists.start_list(ordered=False)
        ordered_b = self.lists.start_list(ordered=True, start=4)
        self.assertEqual(len(self.numbering.findall(qn("w:abstractNum"))), before + 2)
        self.assertIs(self._abstract_for(ordered_a), self._abstract_for(ordered_b))
        self.assertIsNot(self._abstract_for(ordered_a), self._abstract_for(bullet))
        tags = [child.tag for child in self.numbering]
        last_abstract = max(i for i, tag in enumerate(tags) if tag == qn("w:abstractNum"))
        first_num = min(i for i, tag in enumerate(tags) if tag == qn("w:num"))
        self.assertLess(last_abstract, first_num)

    def test_ordered_list_restarts_on_every_level(self) -> None:
        num = self._num(self.lists.start_list(ordered=True, start=7))
        overrides = num.findall(qn("w:lvlOverride"))
        self.assertEqual([o.get(qn("w:ilvl")) for o in overrides], [str(i) for i in range(9)])
        for override in overrides:
            self.assertEqual(override.find(qn("w:startOverride")).get(qn("w:val")), "7")

    def test_decimal_levels_number_with_their_own_counter(self) -> None:
        abstract = self._abstract_for(self.lists.start_list(ordered=True))
        self.assertEqual(abstract.find(qn("w:multiLevelType")).get(qn("w:val")), "hybridMultilevel")
        levels = abstract.findall(qn("w:lvl"))
        self.assertEqual(len(levels), 9)
        for index, level in enumerate(levels):
            self.assertEqual(level.find(qn("w:numFmt")).get(qn("w:val")), "decimal")
            self.assertEqual(level.find(qn("w:lvlText")).get(qn("w:val")), f"%{index + 1}.")

    def test_bullets_cycle_unicode_symbols_without_symbol_font(self) -> None:
        num_id = self.lists.start_list(ordered=False)
        self.assertEqual(self._num(num_id).findall(qn("w:lvlOverride")), [])
        levels = self._abstract_for(num_id).findall(qn("w:lvl"))
        texts = [level.find(qn("w:lvlText")).get(qn("w:val")) for level in levels]
        self.assertEqual(texts[:4], ["•", "◦", "▪", "•"])
        for level in levels:
            self.assertEqual(level.find(qn("w:numFmt")).get(qn("w:val")), "bullet")
            fonts = level.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
            self.assertNotEqual(fonts.get(qn("w:ascii")), "Symbol")

    def test_bullet_font_override(self) -> None:
        lists = ListNumbering(self.document, bullet_font="Arial")
        level = self._abstract_for(lists.start_list(ordered=False)).find(qn("w:lvl"))
        self.assertEqual(level.find(f"{qn('w:rPr')}/{qn('w:rFonts')}").get(qn("w:ascii")), "Arial")

    def test_apply_sets_num_pr_in_schema_order_and_clamps_level(self) -> None:
        paragraph = self.document.add_paragraph("item", style="List Paragraph")
        paragraph.paragraph_format.space_after = Pt(0)
        num_id = self.lists.start_list(ordered=False)
        self.lists.apply(paragraph, num_id, 12)
        num_pr = paragraph._p.pPr.numPr
        self.assertEqual(num_pr.ilvl.val, 8)
        self.assertEqual(num_pr.numId.val, num_id)
        self.assert_schema_order(paragraph._p.pPr, ooxml._PPR_SEQUENCE)
        self.lists.apply(paragraph, num_id, -3)
        self.assertEqual(paragraph._p.pPr.numPr.ilvl.val, 0)

    def test_continuation_indent_matches_item_text_indent(self) -> None:
        levels = self._abstract_for(self.lists.start_list(ordered=True)).findall(qn("w:lvl"))
        for index, level in enumerate(levels):
            left = int(level.find(f"{qn('w:pPr')}/{qn('w:ind')}").get(qn("w:left")))
            self.assertEqual(self.lists.continuation_indent(index), Twips(left))
        self.assertEqual(self.lists.continuation_indent(0), Twips(720))
        self.assertEqual(self.lists.continuation_indent(20), Twips(720 * 9))

    def test_creates_numbering_part_when_document_has_none(self) -> None:
        document = Document()
        rels = document.part.rels
        for r_id in [k for k, rel in rels.items() if rel.reltype == RT.NUMBERING]:
            del rels[r_id]
        lists = ListNumbering(document)
        num_id = lists.start_list(ordered=True)
        lists.apply(document.add_paragraph("one"), num_id, 0)
        reopened, blob = _reopen(document)
        self.assertIn('w:numId="1"', _zip_text(blob, "word/numbering.xml"))
        self.assertEqual(len(reopened.part.numbering_part.element.findall(qn("w:num"))), 1)


class FootnoteTests(_OrderAssertions):
    def setUp(self) -> None:
        self.document = Document()
        self.manager = FootnoteManager(self.document)

    def test_manager_is_lazy(self) -> None:
        with self.assertRaises(KeyError):
            self.document.part.part_related_by(RT.FOOTNOTES)
        self.assertIsNone(self.document.settings.element.find(qn("w:footnotePr")))

    def test_first_reserve_creates_part_settings_and_styles(self) -> None:
        self.assertEqual(self.manager.reserve(), 1)
        part = self.document.part.part_related_by(RT.FOOTNOTES)
        self.assertIsInstance(part, FootnotesPart)
        self.assertEqual(part.partname, "/word/footnotes.xml")
        self.assertEqual(part.content_type, CT.WML_FOOTNOTES)
        separators = {fn.get(qn("w:id")): fn.get(qn("w:type")) for fn in part.element}
        self.assertEqual(separators["-1"], "separator")
        self.assertEqual(separators["0"], "continuationSeparator")

        settings = self.document.settings.element
        footnote_pr = settings.find(qn("w:footnotePr"))
        self.assertEqual([fn.get(qn("w:id")) for fn in footnote_pr], ["-1", "0"])
        self.assert_schema_order(settings, ooxml._SETTINGS_SEQUENCE)

        styles = self.document.styles.element
        text = styles.xpath('./w:style[@w:styleId="FootnoteText"]')[0]
        self.assertEqual(text.get(qn("w:type")), "paragraph")
        self.assertEqual(text.find(qn("w:basedOn")).get(qn("w:val")), "Normal")
        self.assertEqual(text.xpath("./w:rPr/w:sz/@w:val"), ["20"])
        self.assertEqual(text.xpath("./w:pPr/w:spacing/@w:after"), ["0"])
        self.assertEqual(text.xpath("./w:pPr/w:ind/@w:firstLine"), ["0"])
        reference = styles.xpath('./w:style[@w:styleId="FootnoteReference"]')[0]
        self.assertEqual(reference.get(qn("w:type")), "character")
        self.assertEqual(reference.xpath("./w:rPr/w:vertAlign/@w:val"), ["superscript"])

    def test_reserve_returns_increasing_ids(self) -> None:
        self.assertEqual([self.manager.reserve() for _ in range(3)], [1, 2, 3])

    def test_reference_run_is_styled(self) -> None:
        paragraph = self.document.add_paragraph("text")
        footnote_id = self.manager.insert_reference(paragraph)
        run = paragraph._p.findall(qn("w:r"))[-1]
        self.assertEqual(run.xpath("./w:rPr/w:rStyle/@w:val"), ["FootnoteReference"])
        self.assertEqual(run.xpath("./w:footnoteReference/@w:id"), [str(footnote_id)])

    def test_first_paragraph_starts_with_mark_and_space(self) -> None:
        footnote_id = self.manager.reserve()
        self.assertFalse(self.manager.has_content(footnote_id))
        first = self.manager.add_paragraph(footnote_id)
        first.add_run("Body")
        second = self.manager.add_paragraph(footnote_id)
        second.add_run("More")
        self.assertTrue(self.manager.has_content(footnote_id))
        runs = first._p.findall(qn("w:r"))
        self.assertIsNotNone(runs[0].find(qn("w:footnoteRef")))
        self.assertEqual(runs[0].xpath("./w:rPr/w:rStyle/@w:val"), ["FootnoteReference"])
        self.assertEqual(runs[1].xpath("./w:t/text()"), [" "])
        self.assertEqual(first.style.name, "footnote text")
        self.assertEqual(second._p.xpath(".//w:footnoteRef"), [])
        self.assertEqual(second.text, "More")

    def test_paragraph_part_is_the_footnotes_part(self) -> None:
        footnote_id = self.manager.reserve()
        paragraph = self.manager.add_paragraph(footnote_id)
        self.assertIs(paragraph.part, self.manager.part)
        add_hyperlink_run(add_external_hyperlink(paragraph, "https://example.org/n"), paragraph, "x")
        self.assertIn("https://example.org/n", [r.target_ref for r in self.manager.part.rels.values()])
        self.assertNotIn(
            "https://example.org/n",
            [r.target_ref for r in self.document.part.rels.values() if r.is_external],
        )

    def test_style_argument(self) -> None:
        footnote_id = self.manager.reserve()
        self.assertIsNone(self.manager.add_paragraph(footnote_id, style=None)._p.pPr)
        self.assertEqual(self.manager.add_paragraph(footnote_id, style="Normal").style.name, "Normal")
        with self.assertRaises(KeyError):
            self.manager.add_paragraph(footnote_id, style="No Such Style")
        with self.assertRaises(KeyError):
            self.manager.add_paragraph(99)

    def test_rich_footnote_round_trips(self) -> None:
        paragraph = self.document.add_paragraph("Claim")
        footnote_id = self.manager.insert_reference(paragraph)
        body = self.manager.add_paragraph(footnote_id)
        body.add_run("Bold note ").bold = True
        link = add_external_hyperlink(body, "https://example.com/source")
        add_hyperlink_run(link, body, "source")
        shape = body.add_run().add_picture(BytesIO(_image_bytes()), width=Mm(5))
        set_picture_description(shape, "tiny")
        self.assertGreaterEqual(int(shape._inline.docPr.get("id")), 1_000_000)

        reopened, blob = _reopen(self.document)
        xml = _zip_text(blob, "word/footnotes.xml")
        self.assertIn("Bold note", xml)
        self.assertIn('descr="tiny"', xml)
        part = reopened.part.part_related_by(RT.FOOTNOTES)
        self.assertIsInstance(part, FootnotesPart)
        link_id = part.element.xpath(".//w:hyperlink/@r:id")[0]
        self.assertEqual(part.rels[link_id].target_ref, "https://example.com/source")
        blip_id = part.element.xpath(".//a:blip/@r:embed")[0]
        image_part = part.rels[blip_id].target_part
        self.assertIsInstance(image_part, ImagePart)
        self.assertTrue(image_part.blob.startswith(b"\x89PNG"))

        manager = FootnoteManager(reopened)
        self.assertEqual(manager.reserve(), footnote_id + 1)
        self.assertIs(manager.part, part)
        self.assertEqual(manager.text_style_id, "FootnoteText")
        styles = reopened.styles.element
        self.assertEqual(len(styles.xpath('./w:style[@w:styleId="FootnoteText"]')), 1)

    def test_empty_footnote_serializes_with_its_mark_only(self) -> None:
        footnote_id = self.manager.insert_reference(self.document.add_paragraph("x"))
        _, blob = _reopen(self.document)
        footnotes = parse_xml(_zip_text(blob, "word/footnotes.xml").encode("utf-8"))
        saved = footnotes.xpath(f'./w:footnote[@w:id="{footnote_id}"]')[0]
        self.assertEqual(len(saved.xpath("./w:p//w:footnoteRef")), 1)
        self.assertFalse(self.manager.has_content(footnote_id))

    def test_reuses_localized_footnote_styles(self) -> None:
        styles = self.document.styles.element
        styles.append(parse_xml(
            f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="a5">'
            '<w:name w:val="footnote text"/></w:style>'
        ))
        styles.append(parse_xml(
            f'<w:style {nsdecls("w")} w:type="character" w:styleId="a7">'
            '<w:name w:val="footnote reference"/></w:style>'
        ))
        footnote_id = self.manager.insert_reference(self.document.add_paragraph("x"))
        paragraph = self.manager.add_paragraph(footnote_id)
        self.assertEqual(paragraph._p.pPr.pStyle.val, "a5")
        self.assertEqual(paragraph._p.xpath("./w:r/w:rPr/w:rStyle/@w:val"), ["a7"])
        self.assertEqual(styles.xpath('./w:style[@w:styleId="FootnoteText"]'), [])

    def test_upgrades_part_loaded_before_registration(self) -> None:
        footnote_id = self.manager.insert_reference(self.document.add_paragraph("x"))
        note = self.manager.add_paragraph(footnote_id)
        note.add_run("old ")
        note.add_run().add_picture(BytesIO(_image_bytes()))
        stream = BytesIO()
        self.document.save(stream)
        registry = {k: v for k, v in PartFactory.part_type_for.items() if k != CT.WML_FOOTNOTES}
        with patch.dict(PartFactory.part_type_for, registry, clear=True):
            stream.seek(0)
            legacy = Document(stream)
        self.assertNotIsInstance(legacy.part.part_related_by(RT.FOOTNOTES), FootnotesPart)
        manager = FootnoteManager(legacy)
        new_id = manager.reserve()
        manager.add_paragraph(new_id).add_run("new")
        _, blob = _reopen(legacy)
        xml = _zip_text(blob, "word/footnotes.xml")
        self.assertIn("old", xml)
        self.assertIn("new", xml)
        self.assertIn("media/image1.png", _zip_text(blob, "word/_rels/footnotes.xml.rels"))


class SlugAndBookmarkTests(unittest.TestCase):
    def test_github_slug(self) -> None:
        used: dict[str, int] = {}
        self.assertEqual(github_slug("Таблицы", used), "таблицы")
        self.assertEqual(github_slug("Hello, World!", used), "hello-world")
        self.assertEqual(github_slug("Hello", used), "hello")
        self.assertEqual(github_slug("Hello", used), "hello-1")
        self.assertEqual(github_slug("Hello-1", used), "hello-1-1")
        self.assertEqual(github_slug("Hello", used), "hello-2")
        self.assertEqual(github_slug("C++ & Python 3.12", used), "c--python-312")
        self.assertEqual(github_slug("🎉 Party time", used), "-party-time")
        self.assertEqual(github_slug("snake_case and-dash", used), "snake_case-and-dash")
        self.assertEqual(github_slug("  Привет, мир!  ", used), "привет-мир")

    def test_github_slug_skips_slugs_taken_by_real_headings(self) -> None:
        used: dict[str, int] = {}
        self.assertEqual(
            [github_slug(text, used) for text in ("A", "A-1", "A")], ["a", "a-1", "a-2"]
        )

    def test_bookmark_name(self) -> None:
        used: set[str] = set()
        self.assertEqual(bookmark_name("таблицы-и-списки", used), "таблицы_и_списки")
        self.assertEqual(bookmark_name("1-intro", used), "h_1_intro")
        self.assertEqual(bookmark_name("-party", used), "h__party")
        self.assertEqual(bookmark_name("", used), "h_")
        self.assertEqual(bookmark_name("Таблицы-и-списки", used), "Таблицы_и_списки_2")
        self.assertIn("h_1_intro", used)

    def test_bookmark_name_truncates_and_stays_unique(self) -> None:
        used: set[str] = set()
        slug = "очень-длинный-заголовок-" * 4
        names = [bookmark_name(slug, used) for _ in range(12)]
        self.assertEqual(len(set(names)), 12)
        for name in names:
            self.assertLessEqual(len(name), 40)
            self.assertTrue(name[0].isalpha())
        self.assertTrue(names[1].endswith("_2"))
        self.assertTrue(names[11].endswith("_12"))

    def test_add_bookmark_wraps_content_with_unique_ids(self) -> None:
        document = Document()
        existing = document.add_paragraph("pre")
        existing._p.append(parse_xml(
            f'<w:bookmarkStart {nsdecls("w")} w:id="7" w:name="legacy"/>'
        ))
        heading = document.add_heading("Title", level=1)
        first_id = add_bookmark(heading, "title")
        second_id = add_bookmark(document.add_paragraph("other"), "other")
        self.assertEqual((first_id, second_id), (8, 9))
        children = [child.tag for child in heading._p]
        self.assertEqual(children[0], qn("w:pPr"))
        self.assertEqual(children[1], qn("w:bookmarkStart"))
        self.assertEqual(children[-1], qn("w:bookmarkEnd"))
        self.assertIn(qn("w:r"), children[2:-1])
        with self.assertRaises(ValueError):
            add_bookmark(heading, "x" * 41)

    def test_internal_and_external_hyperlinks(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("See ")
        internal = add_internal_hyperlink(paragraph, "таблицы")
        run = add_hyperlink_run(internal, paragraph, "tables")
        self.assertIs(paragraph._p[-1], internal)
        self.assertEqual(internal.get(qn("w:anchor")), "таблицы")
        self.assertEqual(internal.get(qn("w:history")), "1")
        self.assertIs(run._r.getparent(), internal)
        self.assertEqual(run.text, "tables")
        external = add_external_hyperlink(paragraph, "https://example.com/a?b=1&c=2")
        rel = document.part.rels[external.get(qn("r:id"))]
        self.assertTrue(rel.is_external)
        self.assertEqual(rel.target_ref, "https://example.com/a?b=1&c=2")
        self.assertIn("tables", paragraph.text)


class FieldTests(_OrderAssertions):
    def test_add_field_structure(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        add_field(paragraph, "SEQ Figure \\* ARABIC", "1")
        runs = paragraph._p.findall(qn("w:r"))
        kinds = [r.xpath("./w:fldChar/@w:fldCharType") for r in runs]
        self.assertEqual(kinds, [["begin"], [], ["separate"], [], ["end"]])
        self.assertEqual(runs[0].xpath("./w:fldChar/@w:dirty"), ["true"])
        instr = runs[1].find(qn("w:instrText"))
        self.assertEqual(instr.text, " SEQ Figure \\* ARABIC ")
        self.assertEqual(instr.get(qn("xml:space")), "preserve")
        self.assertEqual(runs[3].xpath("./w:t/text()"), ["1"])
        clean = document.add_paragraph()
        self.assertIsNone(add_field(clean, "PAGE", dirty=False))
        self.assertEqual(clean._p.xpath(".//w:fldChar/@w:dirty"), [])
        self.assertEqual(len(clean._p.findall(qn("w:r"))), 4)

    def test_update_fields_setting_is_ordered_and_idempotent(self) -> None:
        document = Document()
        FootnoteManager(document).reserve()
        request_field_update_on_open(document)
        request_field_update_on_open(document)
        settings = document.settings.element
        self.assertEqual(len(settings.findall(qn("w:updateFields"))), 1)
        self.assertEqual(settings.find(qn("w:updateFields")).get(qn("w:val")), "true")
        self.assert_schema_order(settings, ooxml._SETTINGS_SEQUENCE)
        self.assertEqual(settings[-1].tag, "{http://schemas.microsoft.com/office/word/2010/wordml}defaultImageDpi")

    def test_insert_toc_before(self) -> None:
        document = Document()
        anchor = document.add_paragraph("Body")
        toc = insert_toc_before(anchor, "1-2", "Update me", title="Contents")
        paragraphs = document.paragraphs
        self.assertEqual([p.text for p in paragraphs], ["Contents", "Update me", "Body"])
        self.assertEqual(paragraphs[0].style.style_id, "TOCHeading")
        self.assertIs(toc._p, paragraphs[1]._p)
        self.assertIn('TOC \\o "1-2" \\h \\z \\u', toc._p.xpath("string(.//w:instrText)"))
        with self.assertRaises(ValueError):
            insert_toc_before(anchor, "1-x")

    def test_append_toc_falls_back_to_bold_title(self) -> None:
        document = Document()
        styles = document.styles.element
        styles.remove(styles.xpath('./w:style[@w:styleId="TOCHeading"]')[0])
        toc = append_toc(document, title="Оглавление")
        title = document.paragraphs[-2]
        self.assertTrue(title.runs[0].bold)
        self.assertIs(document.paragraphs[-1]._p, toc._p)
        self.assertIn("TOC", toc._p.xpath("string(.//w:instrText)"))

    def test_add_toc_dispatches(self) -> None:
        document = Document()
        anchor = document.add_paragraph("Body")
        add_toc(anchor)
        add_toc(document, levels="1-1")
        self.assertEqual(document.paragraphs[1].text, "Body")
        self.assertIn('"1-1"', document.paragraphs[-1]._p.xpath("string(.//w:instrText)"))


class PageAndLanguageTests(_OrderAssertions):
    def test_apply_page_setup(self) -> None:
        document = Document()
        section = document.sections[0]
        apply_page_setup(section, "a4", (20, 15, 25, 30))
        sect_pr = section._sectPr
        self.assertEqual(sect_pr.xpath("./w:pgSz/@w:w"), ["11906"])
        self.assertEqual(sect_pr.xpath("./w:pgSz/@w:h"), ["16838"])
        # Word stores lengths in twips, so millimetres come back rounded.
        self.assertAlmostEqual(section.left_margin, Mm(30), delta=635)
        self.assertAlmostEqual(section.top_margin, Mm(20), delta=635)
        self.assertAlmostEqual(text_width(section), Mm(210 - 30 - 15), delta=3 * 635)
        self.assertAlmostEqual(text_height(section), Mm(297 - 20 - 25), delta=3 * 635)
        self.assertEqual(
            text_width(section), section.page_width - section.left_margin - section.right_margin
        )
        apply_page_setup(section, "Letter")
        self.assertEqual(sect_pr.xpath("./w:pgSz/@w:w"), ["12240"])
        self.assertEqual(sect_pr.xpath("./w:pgSz/@w:h"), ["15840"])
        with self.assertRaises(ValueError):
            apply_page_setup(section, "A5")
        with self.assertRaises(ValueError):
            apply_page_setup(section, "A4", (10, 10, -1, 10))

    def test_set_document_language(self) -> None:
        document = Document()
        set_document_language(document, "ru-RU")
        lang = document.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:lang")[0]
        self.assertEqual(lang.get(qn("w:val")), "ru-RU")
        self.assertEqual(lang.get(qn("w:eastAsia")), "ru-RU")  # was en-US
        self.assertEqual(lang.get(qn("w:bidi")), "ar-SA")
        self.assertEqual(document.core_properties.language, "ru-RU")
        lang.set(qn("w:eastAsia"), "ja-JP")
        set_document_language(document, "en-GB")
        self.assertEqual(lang.get(qn("w:val")), "en-GB")
        self.assertEqual(lang.get(qn("w:eastAsia")), "ja-JP")
        with self.assertRaises(ValueError):
            set_document_language(document, "not a tag")

    def test_set_document_language_creates_doc_defaults(self) -> None:
        document = Document()
        styles = document.styles.element
        styles.remove(styles.find(qn("w:docDefaults")))
        set_document_language(document, "ru-RU")
        self.assertEqual(styles[0].tag, qn("w:docDefaults"))
        self.assertEqual(styles.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:lang/@w:val"), ["ru-RU"])

    def test_detect_language(self) -> None:
        self.assertEqual(detect_language(""), "en-US")
        self.assertEqual(detect_language("123 !!!"), "en-US")
        self.assertEqual(detect_language("Привет, мир"), "ru-RU")
        self.assertEqual(detect_language("abcdefg абв"), "ru-RU")  # exactly 30% Cyrillic
        self.assertEqual(detect_language("abcdefgh абв"), "en-US")  # 27%
        self.assertEqual(detect_language("Hello wonderful world, да"), "en-US")

    def test_add_page_number_footer(self) -> None:
        document = Document()
        section = document.sections[0]
        paragraph = add_page_number_footer(section, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertFalse(section.footer.is_linked_to_previous)
        self.assertEqual(paragraph.alignment, WD_ALIGN_PARAGRAPH.RIGHT)
        self.assertEqual(paragraph._p.xpath("string(.//w:instrText)").strip(), "PAGE")
        self.assertEqual(len(section.footer.paragraphs), 1)


class ImageTests(unittest.TestCase):
    def test_picture_size_to_fit(self) -> None:
        big = Inches(100)
        self.assertEqual(picture_size_to_fit((960, 480), (96, 96), big, big), (Inches(10), Inches(5)))
        self.assertEqual(picture_size_to_fit((960, 480), (0, 0), big, big), (Inches(10), Inches(5)))
        self.assertEqual(picture_size_to_fit((960, 480), None, big, big), (Inches(10), Inches(5)))
        self.assertEqual(picture_size_to_fit((600, 300), (300, 300), big, big), (Inches(2), Inches(1)))
        width, height = picture_size_to_fit((960, 480), (96, 96), Inches(6), big)
        self.assertEqual((width, height), (Inches(6), Inches(3)))
        width, height = picture_size_to_fit((480, 960), (96, 96), Inches(6), Inches(4))
        self.assertEqual((width, height), (Inches(2), Inches(4)))
        self.assertEqual(picture_size_to_fit((10, 10), (96, 96), Inches(6), Inches(6)), (Emu(95250), Emu(95250)))
        with self.assertRaises(ValueError):
            picture_size_to_fit((0, 10), (96, 96), big, big)

    def test_svg_intrinsic_size(self) -> None:
        def size(attrs: str):
            return svg_intrinsic_size(f'<svg xmlns="http://www.w3.org/2000/svg" {attrs}/>'.encode())

        self.assertEqual(size('width="100" height="50"'), (100.0, 50.0))
        self.assertEqual(size('width="2in" height="72pt"'), (192.0, 96.0))
        self.assertAlmostEqual(size('width="25.4mm" height="2.54cm"')[0], 96.0)
        self.assertEqual(size('viewBox="0 0 300 150"'), (300.0, 150.0))
        self.assertEqual(size('width="600" viewBox="0,0,300,150"'), (600.0, 300.0))
        self.assertEqual(size('height="10" viewBox="0 0 300 150"'), (20.0, 10.0))
        self.assertEqual(size('width="100%" height="100%" viewBox="0 0 40 20"'), (40.0, 20.0))
        self.assertIsNone(size('width="100%"'))
        self.assertIsNone(svg_intrinsic_size(b"<html/>"))
        self.assertIsNone(svg_intrinsic_size(b"not xml"))

    def test_svg_parsing_does_not_expand_entities(self) -> None:
        hostile = (
            b'<?xml version="1.0"?><!DOCTYPE svg [<!ENTITY a "aaaaaaaaaa">'
            b'<!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">'
            b'<!ENTITY xxe SYSTEM "file:///etc/passwd">]>'
            b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="5"><text>&b;&xxe;</text></svg>'
        )
        root = ooxml._parse_svg_root(hostile)
        self.assertNotIn("aaaaaaaaaa", "".join(root.itertext()))
        self.assertNotIn("root:", "".join(root.itertext()))
        self.assertEqual(svg_intrinsic_size(hostile), (10.0, 5.0))

    def test_add_svg_picture(self) -> None:
        document = Document()
        run = document.add_paragraph().add_run()
        shape = add_svg_picture(run, _SVG)
        self.assertEqual((shape.width, shape.height), (Emu(200 * 9525), Emu(100 * 9525)))
        blip = shape._inline.xpath(".//a:blip")[0]
        ext = blip.xpath("./a:extLst/a:ext")[0]
        self.assertEqual(ext.get("uri"), "{96DAC541-7B7A-43D3-8B79-37D633B846F1}")
        svg_blip = ext.find(f"{_ASVG}svgBlip")
        svg_part = document.part.rels[svg_blip.get(qn("r:embed"))].target_part
        self.assertEqual(svg_part.content_type, "image/svg+xml")
        self.assertEqual(svg_part.blob, ooxml.sanitize_svg(_SVG))
        self.assertEqual(document.part.rels[svg_blip.get(qn("r:embed"))].reltype, RT.IMAGE)
        fallback = document.part.rels[blip.get(qn("r:embed"))].target_part
        with Image.open(BytesIO(fallback.blob)) as image:
            self.assertEqual(image.format, "PNG")
            self.assertEqual(image.size, (256, 128))

        second = add_svg_picture(document.add_paragraph().add_run(), _SVG, width=Inches(1))
        self.assertEqual((second.width, second.height), (Inches(1), Inches(0.5)))
        svg_parts = [p for p in document.part.package.image_parts if p.content_type == "image/svg+xml"]
        self.assertEqual(len(svg_parts), 1)

        custom = _image_bytes(size=(8, 8))
        third = add_svg_picture(document.add_paragraph().add_run(), _SVG, height=Inches(2), fallback_png=custom)
        self.assertEqual((third.width, third.height), (Inches(4), Inches(2)))
        third_blip = third._inline.xpath(".//a:blip")[0]
        self.assertEqual(document.part.rels[third_blip.get(qn("r:embed"))].target_part.blob, custom)

        reopened, blob = _reopen(document)
        self.assertIn("image/svg+xml", _zip_text(blob, "[Content_Types].xml"))
        self.assertEqual(len(reopened.inline_shapes), 3)
        with self.assertRaises(ValueError):
            add_svg_picture(document.add_paragraph().add_run(), b"<html/>")

    def test_sanitize_svg(self) -> None:
        cleaned = ooxml.sanitize_svg(_ILLUSTRATOR_SVG)
        for forbidden in (b"<!DOCTYPE", b"ENTITY", b"&ns_svg;", b"&st0;", b"<!--"):
            self.assertNotIn(forbidden, cleaned)
        root = parse_xml(cleaned)
        self.assertEqual(root.tag, "{http://www.w3.org/2000/svg}svg")
        self.assertEqual(root.getroottree().docinfo.doctype, "")
        self.assertEqual(root.find("{http://www.w3.org/2000/svg}rect").get("style"), "fill:#3a7")
        self.assertEqual(root.find("{http://www.w3.org/2000/svg}text").text, "a & b <x>")
        self.assertTrue(cleaned.startswith(b"<?xml version='1.0' encoding='UTF-8'?>"))
        self.assertEqual(parse_xml(ooxml.sanitize_svg(_SVG)).get("width"), "200")
        with self.assertRaises(ValueError):
            ooxml.sanitize_svg(b"<html/>")

    def test_sanitize_svg_refuses_external_and_exploding_entities(self) -> None:
        # External entity: never fetched; the picture degrades to the PNG alone.
        self.assertIsNone(ooxml.sanitize_svg(_EXTERNAL_ENTITY_SVG))
        document = Document()
        shape = add_svg_picture(document.add_paragraph().add_run(), _EXTERNAL_ENTITY_SVG)
        self.assertEqual(shape._inline.xpath(".//a:blip/a:extLst"), [])
        self.assertEqual((shape.width, shape.height), (Emu(10 * 9525), Emu(5 * 9525)))
        # Entity amplification: libxml2 refuses it outright -> not a usable SVG.
        with self.assertRaisesRegex(ValueError, "invalid SVG"):
            ooxml.sanitize_svg(_BILLION_LAUGHS_SVG)
        with self.assertRaisesRegex(ValueError, "invalid SVG"):
            add_svg_picture(Document().add_paragraph().add_run(), _BILLION_LAUGHS_SVG)

    def test_svg_with_doctype_is_stored_sanitized(self) -> None:
        document = Document()
        shape = add_svg_picture(document.add_paragraph().add_run(), _ILLUSTRATOR_SVG)
        self.assertEqual((shape.width, shape.height), (Emu(120 * 9525), Emu(60 * 9525)))
        r_id = shape._inline.xpath(".//a:blip/a:extLst/a:ext")[0].find(f"{_ASVG}svgBlip").get(qn("r:embed"))
        stored = document.part.rels[r_id].target_part.blob
        self.assertNotIn(b"<!DOCTYPE", stored)
        self.assertEqual(stored, ooxml.sanitize_svg(_ILLUSTRATOR_SVG))

    def test_svg_falls_back_to_png_only_when_sanitizing_fails(self) -> None:
        document = Document()
        # add_svg_picture looks sanitize_svg up in its own module, ooxml.svg.
        with patch.object(ooxml.svg, "sanitize_svg", return_value=None):
            shape = add_svg_picture(document.add_paragraph().add_run(), _SVG)
        self.assertEqual(shape._inline.xpath(".//a:blip/a:extLst"), [])
        self.assertEqual(
            [p for p in document.part.package.image_parts if p.content_type == "image/svg+xml"], []
        )
        self.assertEqual((shape.width, shape.height), (Emu(200 * 9525), Emu(100 * 9525)))

    def test_set_picture_description(self) -> None:
        document = Document()
        run = document.add_paragraph().add_run()
        shape = run.add_picture(BytesIO(_image_bytes()))
        set_picture_description(shape, "A chart", title="Chart")
        doc_pr = shape._inline.docPr
        self.assertEqual((doc_pr.get("descr"), doc_pr.get("title")), ("A chart", "Chart"))
        set_picture_description(run, "Updated")
        self.assertEqual(doc_pr.get("descr"), "Updated")
        with self.assertRaises(ValueError):
            set_picture_description(document.add_paragraph().add_run(), "none")

    def test_normalize_raster(self) -> None:
        for fmt, expected in (("PNG", "png"), ("JPEG", "jpeg"), ("GIF", "gif"), ("BMP", "bmp")):
            data = _image_bytes(fmt)
            self.assertEqual(normalize_raster(data), (data, expected), fmt)
        for fmt, mode in (("WEBP", "RGBA"), ("ICO", "RGBA")):
            data, kind = normalize_raster(_image_bytes(fmt, size=(16, 16), mode=mode))
            self.assertEqual(kind, "png", fmt)
            with Image.open(BytesIO(data)) as image:
                self.assertEqual((image.format, image.size), ("PNG", (16, 16)))
            Document().add_paragraph().add_run().add_picture(BytesIO(data))
        with self.assertRaisesRegex(ValueError, "unsupported image format"):
            normalize_raster(b"definitely not an image")


class FormattingTests(_OrderAssertions):
    def setUp(self) -> None:
        self.document = Document()

    def test_paragraph_shading_and_borders_in_schema_order(self) -> None:
        paragraph = self.document.add_paragraph("code", style="List Paragraph")
        paragraph.paragraph_format.keep_with_next = True
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        ListNumbering(self.document).apply(paragraph, 1, 0)
        set_paragraph_shading(paragraph, "#f2f2f2")
        set_paragraph_borders(
            paragraph, left=("single", 18, 4, "a0a0a0"), bottom=("single", 4, 1, "DDDDDD")
        )
        set_paragraph_shading(paragraph, "EEEEEE")
        ppr = paragraph._p.pPr
        self.assertEqual(ppr.xpath("./w:shd/@w:fill"), ["EEEEEE"])
        self.assertEqual(ppr.xpath("./w:shd/@w:val"), ["clear"])
        self.assertEqual([child.tag for child in ppr.find(qn("w:pBdr"))], [qn("w:left"), qn("w:bottom")])
        left = ppr.find(qn("w:pBdr")).find(qn("w:left"))
        self.assertEqual(
            [left.get(qn(f"w:{a}")) for a in ("val", "sz", "space", "color")],
            ["single", "18", "4", "A0A0A0"],
        )
        self.assert_schema_order(ppr, ooxml._PPR_SEQUENCE)
        self.assert_schema_order(ppr.find(qn("w:pBdr")), ooxml._PBDR_SEQUENCE)
        set_paragraph_borders(paragraph)
        self.assertIsNone(ppr.find(qn("w:pBdr")))
        with self.assertRaises(ValueError):
            set_paragraph_shading(paragraph, "grey")

    def test_style_targets(self) -> None:
        quote = self.document.styles["Quote"]
        set_paragraph_shading(quote, "FAFAFA")
        set_paragraph_borders(quote, left=("single", 24, 8, "C0C0C0"))
        set_keep_lines(quote, True, True)
        ppr = quote.element.pPr
        self.assertEqual(ppr.xpath("./w:shd/@w:fill"), ["FAFAFA"])
        self.assertTrue(quote.paragraph_format.keep_together)
        self.assertTrue(quote.paragraph_format.keep_with_next)
        self.assert_schema_order(ppr, ooxml._PPR_SEQUENCE)
        strong = self.document.styles["Strong"]
        set_run_shading(strong, "EEEEEE")
        self.assertEqual(strong.element.rPr.xpath("./w:shd/@w:fill"), ["EEEEEE"])

    def test_run_shading_in_schema_order(self) -> None:
        run = self.document.add_paragraph().add_run("x")
        run.bold = True
        run.font.superscript = True
        run.font.size = Pt(9)
        set_run_shading(run, "eeeeee")
        rpr = run._r.rPr
        self.assertEqual(rpr.xpath("./w:shd/@w:fill"), ["EEEEEE"])
        self.assert_schema_order(rpr, ooxml._RPR_SEQUENCE)

    def test_row_flags(self) -> None:
        table = self.document.add_table(rows=2, cols=2)
        row = table.rows[0]
        set_table_header_repeat(row)
        set_cant_split(row)
        set_table_header_repeat(row)
        tr_pr = row._tr.trPr
        self.assertEqual([child.tag for child in tr_pr], [qn("w:cantSplit"), qn("w:tblHeader")])

    def test_keep_lines_on_paragraph(self) -> None:
        paragraph = self.document.add_paragraph("x")
        set_keep_lines(paragraph)
        self.assertTrue(paragraph.paragraph_format.keep_together)
        self.assertIsNone(paragraph.paragraph_format.keep_with_next)
        set_keep_lines(paragraph, False, True)
        self.assertFalse(paragraph.paragraph_format.keep_together)
        self.assertTrue(paragraph.paragraph_format.keep_with_next)


def _build_everything(path: Path) -> None:
    """One document exercising every primitive, for package-level validation."""
    document = Document()
    section = document.sections[0]
    apply_page_setup(section, "A4", (20, 20, 20, 25))
    set_document_language(document, detect_language("Проверка документа"))
    add_page_number_footer(section)
    request_field_update_on_open(document)
    append_toc(document, title="Содержание")
    slugs: dict[str, int] = {}
    names: set[str] = set()
    heading = document.add_heading("Списки и сноски", level=1)
    add_bookmark(heading, bookmark_name(github_slug(heading.text, slugs), names))
    link_paragraph = document.add_paragraph("Go to ")
    add_hyperlink_run(add_internal_hyperlink(link_paragraph, "списки_и_сноски"), link_paragraph, "lists")
    lists = ListNumbering(document)
    outer = lists.start_list(ordered=True, start=2)
    inner = lists.start_list(ordered=False)
    for text, num_id, level in (("one", outer, 0), ("nested", inner, 1), ("two", outer, 0)):
        lists.apply(document.add_paragraph(text), num_id, level)
    document.add_paragraph("continued").paragraph_format.left_indent = lists.continuation_indent(0)
    manager = FootnoteManager(document)
    claim = document.add_paragraph("Claim")
    note = manager.add_paragraph(manager.insert_reference(claim))
    note.add_run("Bold").bold = True
    add_hyperlink_run(add_external_hyperlink(note, "https://example.com"), note, " link")
    note.add_run().add_picture(BytesIO(_image_bytes()), width=Mm(4))
    manager.insert_reference(claim)  # left empty on purpose
    code = document.add_paragraph("code")
    set_paragraph_shading(code, "F2F2F2")
    set_paragraph_borders(code, left=("single", 18, 4, "A0A0A0"))
    set_keep_lines(code, True, True)
    set_run_shading(document.add_paragraph().add_run("inline"), "EEEEEE")
    table = document.add_table(rows=2, cols=2)
    set_table_header_repeat(table.rows[0])
    set_cant_split(table.rows[1])
    add_svg_picture(document.add_paragraph().add_run(), _SVG, width=Mm(30))
    add_svg_picture(document.add_paragraph().add_run(), _ILLUSTRATOR_SVG)
    raster, _kind = normalize_raster(_image_bytes("WEBP", size=(8, 8), mode="RGBA"))
    shape = document.add_paragraph().add_run().add_picture(BytesIO(raster))
    set_picture_description(shape, "converted webp")
    insert_toc_before(code, "1-2")
    document.save(path)


class PackageValidationTests(unittest.TestCase):
    def test_everything_round_trips(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "everything.docx"
            _build_everything(path)
            reopened = Document(path)
            self.assertIsInstance(reopened.part.part_related_by(RT.FOOTNOTES), FootnotesPart)
            self.assertEqual(len(reopened.inline_shapes), 3)

    @unittest.skipUnless(_OFFICECLI, "officecli is not installed")
    def test_officecli_schema_validation_passes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "everything.docx"
            _build_everything(path)
            result = subprocess.run(
                [_OFFICECLI, "validate", str(path)], capture_output=True, text=True, timeout=120
            )
            self.assertEqual(result.returncode, 0, result.stdout + result.stderr)
            self.assertIn("no errors", result.stdout)


if __name__ == "__main__":
    unittest.main()
