"""Regressions found by reviewing and fuzzing the 1.2 renderer.

Each test reproduces one confirmed defect with the smallest input that hit
it. Where the defect was an invalid file, the saved document is also checked
with ``officecli validate`` when officecli is installed.
"""

from io import BytesIO
from pathlib import Path
import shutil
import subprocess
import tempfile
import time
import unittest

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

from mdtoword.converters import MarkdownToWordConverter
from mdtoword.gfm_renderer import GfmDocxRenderer
from mdtoword.md_extensions import parse_front_matter
from mdtoword.options import DocumentOptions

_OFFICECLI = shutil.which("officecli")
_MATH_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"


def _render(markdown, **options):
    renderer = GfmDocxRenderer("Times New Roman", Pt(12), document_options=DocumentOptions(**options))
    return renderer.render(markdown)


def _png():
    buffer = BytesIO()
    Image.new("RGB", (4, 4), (0, 0, 0)).save(buffer, format="PNG")
    return buffer.getvalue()


class _Validates(unittest.TestCase):
    def assertValidDocx(self, document):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "out.docx"
            document.save(str(path))
            Document(str(path))  # python-docx can read it back
            if _OFFICECLI is None:
                return
            result = subprocess.run(
                [_OFFICECLI, "validate", str(path)], capture_output=True, text=True, timeout=120
            )
            self.assertIn("no errors", result.stdout + result.stderr, result.stdout + result.stderr)


class InputRobustnessTests(_Validates):
    def test_long_front_matter_values_are_shortened_not_fatal(self):
        abstract = "слово " * 80
        document, warnings = _render(f"---\ntitle: T\nabstract: >\n  {abstract}\n---\n# T\n")
        self.assertLessEqual(len(document.core_properties.comments), 255)
        self.assertIn("front_matter_truncated", [w.code for w in warnings])

    def test_control_characters_are_removed_with_one_warning(self):
        document, warnings = _render("```\n\x1b[31mred\x1b[0m\n```\n\ntext\x0cmore\x01\n")
        texts = "\n".join(p.text for p in document.paragraphs)
        self.assertIn("[31mred", texts)
        self.assertNotIn("\x1b", texts)
        self.assertEqual([w.code for w in warnings], ["control_characters_removed"])
        self.assertValidDocx(document)

    def test_control_character_in_html_attribute(self):
        document, _ = _render('<a href="http://x.com/&#1;">x</a>\n')
        self.assertValidDocx(document)

    def test_invalid_front_matter_language_falls_back(self):
        document, warnings = _render("---\nlang: Deutsch\n---\nHallo Welt.\n")
        self.assertIn("front_matter_ignored", [w.code for w in warnings])
        _, warnings = _render("---\nlang: fr_FR\n---\nBonjour.\n")
        self.assertEqual(warnings, [])

    def test_byte_order_mark_does_not_hide_the_front_matter(self):
        document, _ = _render("\ufeff---\ntitle: X\n---\n# Title\n")
        self.assertEqual(document.paragraphs[0].style.name, "Title")
        self.assertEqual(document.paragraphs[1].style.name, "Heading 1")

    def test_byte_order_mark_in_a_file(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "bom.md"
            source.write_bytes("\ufeff# Title\n".encode("utf-8"))
            output = Path(directory) / "bom.docx"
            MarkdownToWordConverter().convert_file(source, output)
            self.assertEqual(Document(str(output)).paragraphs[0].style.name, "Heading 1")

    def test_quoted_table_at_end_of_input_without_newline(self):
        document, _ = _render("> | a |\n> |---|\n>")
        self.assertValidDocx(document)


class FootnoteTests(_Validates):
    def test_footnote_inside_a_footnote_keeps_the_file_valid(self):
        document, warnings = _render("A[^a].\n\n[^a]: see[^b]\n[^b]: inner\n")
        footnotes = next(
            rel.target_part.blob.decode()
            for rel in document.part.rels.values()
            if rel.reltype.endswith("/footnotes")
        )
        self.assertNotIn("footnoteReference", footnotes)
        self.assertIn("inner", footnotes)
        self.assertIn("footnote_nested", [w.code for w in warnings])
        self.assertValidDocx(document)

    def test_section_mode_numbers_notes_in_order(self):
        document, _ = _render(
            "Claim[^source] and inline^[note].\n\n[^source]: Where.\n", footnotes="section"
        )
        body = document.paragraphs[0]
        marks = [run.text for run in body.runs if run.font.superscript]
        self.assertEqual(marks, ["1", "2"])

    def test_section_mode_note_starting_with_code_keeps_its_mark(self):
        document, _ = _render("A[^1].\n\n[^1]:\n    ```\n    code\n    ```\n", footnotes="section")
        texts = [p.text for p in document.paragraphs]
        self.assertTrue(any(text.startswith("1 ") for text in texts), texts)

    def test_copied_footnotes_get_unique_picture_ids(self):
        with tempfile.TemporaryDirectory() as directory:
            (Path(directory) / "a.png").write_bytes(_png())
            document, _ = GfmDocxRenderer("Arial", Pt(12)).render(
                "A[^1] B[^1]\n\n[^1]: pic ![x](a.png)\n", source_path=Path(directory) / "d.md"
            )
        part = next(
            rel.target_part.element
            for rel in document.part.rels.values()
            if rel.reltype.endswith("/footnotes")
        )
        ids = [element.get("id") for element in part.iter(qn("wp:docPr"))]
        self.assertEqual(len(ids), 2)
        self.assertEqual(len(set(ids)), 2)
        self.assertValidDocx(document)


class LinkContentTests(unittest.TestCase):
    def test_math_inside_link_text_stays_inside_the_link(self):
        document, _ = _render("[the $x^2$ value](http://x.com)\n")
        link = document.paragraphs[0]._p.find(qn("w:hyperlink"))
        self.assertIsNotNone(link.find(f"{_MATH_NS}oMath"))

    def test_badge_image_is_clickable(self):
        with tempfile.TemporaryDirectory() as directory:
            (Path(directory) / "b.png").write_bytes(_png())
            document, _ = GfmDocxRenderer("Arial", Pt(12)).render(
                "Badge [![b](b.png)](https://x.com) here\n", source_path=Path(directory) / "d.md"
            )
        link = document.paragraphs[0]._p.find(qn("w:hyperlink"))
        self.assertIsNotNone(link.find(f".//{qn('w:drawing')}"))

    def test_inline_double_dollar_has_no_stray_dollars(self):
        document, warnings = _render("a $$\\frac{a}{b}$$ b\n")
        self.assertNotIn("$", document.paragraphs[0].text)
        self.assertEqual(len(list(document.paragraphs[0]._p.iter(f"{_MATH_NS}oMath"))), 1)
        self.assertEqual(warnings, [])


class HtmlBlockTests(unittest.TestCase):
    def test_link_wrapping_several_paragraphs(self):
        document, _ = _render('<a href="http://x.com">\n<p>Card title</p>\n<p>desc</p>\n</a>\n')
        paragraphs = [p for p in document.paragraphs if p.text.strip()]
        self.assertEqual([p.text for p in paragraphs], ["Card title", "desc"])
        for paragraph in paragraphs:
            self.assertIsNotNone(paragraph._p.find(qn("w:hyperlink")))

    def test_comparison_signs_are_text(self):
        document, warnings = _render("<p>Price: 5 < 10 and 7 > 3</p>\n")
        self.assertEqual(document.paragraphs[0].text, "Price: 5 < 10 and 7 > 3")

    def test_self_closing_iframe_does_not_swallow_the_rest(self):
        document, _ = _render('<div><iframe src="x"/><p>Caption text</p></div>\n')
        self.assertIn("Caption text", "\n".join(p.text for p in document.paragraphs))

    def test_stray_closing_noscript_does_not_swallow_the_rest(self):
        document, _ = _render("<div></noscript><p>Kept</p></div>\n")
        self.assertIn("Kept", "\n".join(p.text for p in document.paragraphs))


class TableTests(unittest.TestCase):
    def test_large_table_renders_in_linear_time(self):
        rows = "\n".join(f"| {i} | b | c | d |" for i in range(500))
        start = time.monotonic()
        document, _ = _render("| a | b | c | d |\n|---|---|---|---|\n" + rows + "\n")
        self.assertLess(time.monotonic() - start, 5)
        self.assertEqual(len(document.tables[0].rows), 501)

    def test_many_columns_stay_inside_the_page(self):
        header = "|" + "|".join(f" c{i} " for i in range(14)) + "|"
        separator = "|" + "|".join("---" for _ in range(14)) + "|"
        document, _ = _render(f"{header}\n{separator}\n")
        section = document.sections[0]
        width = sum(column.width for column in document.tables[0].columns)
        self.assertLessEqual(width, section.page_width - section.left_margin - section.right_margin)


class TemplateTests(_Validates):
    def test_template_with_legacy_compat_settings_stays_valid(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Document()
            compat = template.settings.element.find(qn("w:compat"))
            # In CT_Compat order, ahead of the template's own useFELayout.
            for position, name in enumerate(
                ("spaceForUL", "balanceSingleByteDoubleByteWidth", "ulTrailSpace")
            ):
                compat.insert(position, OxmlElement(f"w:{name}"))
            path = Path(directory) / "t.docx"
            template.save(str(path))
            document, _ = _render("x\n", template=path)
            self.assertValidDocx(document)


class HeadingBookmarkTests(unittest.TestCase):
    def test_heading_in_a_skipped_footnote_does_not_shift_bookmarks(self):
        document, _ = _render(
            "# One\n\n[link](#two)\n\n# Two\n\n[^x]: orphan\n\n    # inner\n"
        )
        bookmarks = {
            b.get(qn("w:name")) for b in document.element.body.iter(qn("w:bookmarkStart"))
        }
        anchor = document.paragraphs[1]._p.find(qn("w:hyperlink")).get(qn("w:anchor"))
        self.assertIn(anchor, bookmarks)
        two = [p for p in document.paragraphs if p.text == "Two"][0]
        self.assertEqual(two._p.find(qn("w:bookmarkStart")).get(qn("w:name")), anchor)


class FrontMatterParserTests(unittest.TestCase):
    def test_quoted_value_with_trailing_comment(self):
        data, _ = parse_front_matter("k: 'a' # c\nd: \"b\" # c\n")
        self.assertEqual(data, {"k": "a", "d": "b"})

    def test_double_quoted_escapes_in_one_pass(self):
        data, _ = parse_front_matter('path: "C:\\\\new"\nline: "a\\nb"\n')
        self.assertEqual(data["path"], "C:\\new")
        self.assertEqual(data["line"], "a\nb")

    def test_mapping_inside_a_block_list_is_reported(self):
        data, skipped = parse_front_matter("authors:\n  - name: A\n    affil: B\n")
        self.assertIn("authors", skipped)


if __name__ == "__main__":
    unittest.main()
