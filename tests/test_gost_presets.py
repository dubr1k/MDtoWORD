"""Saved-package regression checks for the two independent GOST variants."""
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.shared import Mm, Pt
from lxml import etree

from mdtoword.converters import MarkdownToWordConverter
from mdtoword.options import DocumentOptions, PRESET_FONT_SIZE

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
SOURCE = 'Текст[^source].\n\n[^source]: Иванов И. И. Название. — Москва : Издательство, 2026. — 12 с.\n'


class GostPresetPackageTests(unittest.TestCase):
    def package(self, preset, **options):
        converter = MarkdownToWordConverter(
            font_name=options.pop('font_name', 'Times New Roman'),
            font_size=Pt(options.pop('font_size', PRESET_FONT_SIZE[preset])),
            document_options=DocumentOptions(preset=preset, **options),
        )
        with tempfile.TemporaryDirectory() as directory:
            target = Path(directory) / 'result.docx'
            self.assertEqual(converter.convert_content(SOURCE, target), [])
            return target.read_bytes()

    def test_both_saved_presets_and_native_automatically_numbered_footnotes(self):
        for preset, size, vertical in [('gost', 14, 20), ('gost_user', 12, 15)]:
            with self.subTest(preset=preset):
                blob = self.package(preset)
                doc = Document(BytesIO(blob))
                for section in doc.sections:
                    for actual, expected in [
                        (section.page_width.mm, 210), (section.page_height.mm, 297),
                        (section.left_margin.mm, 30), (section.right_margin.mm, 15),
                        (section.top_margin.mm, vertical), (section.bottom_margin.mm, vertical),
                    ]:
                        self.assertAlmostEqual(actual, expected, delta=0.05)
                with ZipFile(BytesIO(blob)) as archive:
                    root = etree.fromstring(archive.read('word/document.xml'))
                    styles = etree.fromstring(archive.read('word/styles.xml'))
                    normal = styles.xpath('//w:style[@w:styleId="Normal"]', namespaces=NS)[0]
                    self.assertEqual(normal.xpath('./w:rPr/w:rFonts/@w:ascii', namespaces=NS), ['Times New Roman'])
                    self.assertEqual(normal.xpath('./w:rPr/w:sz/@w:val', namespaces=NS), [str(size * 2)])
                    self.assertEqual(normal.xpath('./w:rPr/w:color/@w:val', namespaces=NS), ['000000'])
                    self.assertEqual(normal.xpath('./w:pPr/w:spacing/@w:line', namespaces=NS), ['360'])
                    self.assertEqual(normal.xpath('./w:pPr/w:spacing/@w:lineRule', namespaces=NS), ['auto'])
                    refs = root.xpath('//w:footnoteReference/@w:id', namespaces=NS)
                    self.assertEqual(len(refs), 1)
                    notes = etree.fromstring(archive.read('word/footnotes.xml'))
                    note = notes.xpath('//w:footnote[@w:id=$id]', namespaces=NS, id=refs[0])[0]
                    self.assertEqual(len(note.xpath('.//w:footnoteRef', namespaces=NS)), 1)
                    self.assertIn('Иванов', ''.join(note.xpath('.//w:t/text()', namespaces=NS)))
                    self.assertFalse(note.xpath('.//w:footnoteReference', namespaces=NS))
                    if preset == 'gost_user':
                        self.assertEqual(normal.xpath('./w:pPr/w:jc/@w:val', namespaces=NS), ['both'])
                        self.assertFalse(root.xpath('//w:sectPr/w:headerReference | //w:sectPr/w:footerReference', namespaces=NS))
                        self.assertFalse(any(n.startswith(('word/header', 'word/footer')) for n in archive.namelist()))
                    else:
                        self.assertTrue(root.xpath('//w:sectPr/w:footerReference', namespaces=NS))
                        self.assertIn('PAGE', archive.read('word/footer1.xml').decode())

    def test_explicit_font_override(self):
        doc = Document(BytesIO(self.package('gost_user', font_name='Arial')))
        self.assertEqual(doc.styles['Normal'].font.name, 'Arial')
        self.assertEqual(doc.paragraphs[0].runs[0].font.name, 'Arial')

    def test_explicit_size_and_page_override(self):
        doc = Document(BytesIO(self.package('gost_user', font_size=16, page_size='Letter')))
        self.assertEqual(doc.styles['Normal'].font.size.pt, 16)
        self.assertAlmostEqual(doc.sections[0].page_width.mm, 215.9, delta=0.05)

    def test_explicit_template_keeps_its_styles_margins_and_headers(self):
        with tempfile.TemporaryDirectory() as directory:
            template = Document()
            template.styles['Normal'].font.name = 'Georgia'
            template.styles['Normal'].font.size = Pt(18)
            template.sections[0].top_margin = Mm(25)
            template.sections[0].header.paragraphs[0].text = 'Template header'
            path = Path(directory) / 'template.docx'
            template.save(path)
            doc = Document(BytesIO(self.package('gost_user', template=path)))
            self.assertEqual(doc.styles['Normal'].font.name, 'Georgia')
            self.assertEqual(doc.styles['Normal'].font.size.pt, 18)
            self.assertAlmostEqual(doc.sections[0].top_margin.mm, 25, delta=0.05)
            self.assertEqual(doc.sections[0].header.paragraphs[0].text, 'Template header')
            self.assertIsNone(doc.paragraphs[0].runs[0].font.name)

    def test_explicit_section_notes_override_is_still_supported(self):
        with ZipFile(BytesIO(self.package('gost_user', footnotes='section'))) as archive:
            self.assertNotIn('word/footnotes.xml', archive.namelist())
            self.assertIn('Иванов', archive.read('word/document.xml').decode())
