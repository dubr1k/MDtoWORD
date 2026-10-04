"""Tests for the Word -> Markdown converter (``mdtoword.docx_to_markdown``).

Every fixture is built with python-docx directly, with raw OXML where
python-docx has no API (numbering, hyperlinks, fields, footnotes, equations,
drawings), so each test pins one behaviour of the converter in isolation.
"""

from __future__ import annotations

import struct
import tempfile
import unittest
import zlib
from io import BytesIO
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import XmlPart
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt
from markdown_it import MarkdownIt

from mdtoword.docx_to_markdown import WordToMarkdownConverter
from mdtoword.errors import ConversionError, ConversionWarning
from mdtoword.latex_omml import latex_to_omml

_WPS = "http://schemas.microsoft.com/office/word/2010/wordprocessingShape"


def _png(rgb: tuple[int, int, int] = (255, 0, 0)) -> bytes:
    """A valid 1x1 PNG, so python-docx can read its header."""

    def chunk(tag: bytes, data: bytes) -> bytes:
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    raw = b"\x00" + bytes(rgb)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header)
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _run(paragraph, text: str, **formatting):
    run = paragraph.add_run(text)
    for name, value in formatting.items():
        if name == "font":
            run.font.name = value
        else:
            setattr(run.font, name, value)
    return run


def _xml_run(text: str, properties: str = "") -> str:
    return (f"<w:r>{f'<w:rPr>{properties}</w:rPr>' if properties else ''}"
            f'<w:t xml:space="preserve">{text}</w:t></w:r>')


def _append_xml(paragraph, xml: str) -> None:
    fragment = parse_xml(f"<w:p {nsdecls('w', 'r', 'm')}>{xml}</w:p>")
    for child in list(fragment):
        paragraph._p.append(child)


def _hyperlink(paragraph, text: str, url: str | None = None, anchor: str | None = None,
               properties: str = "") -> None:
    link = OxmlElement("w:hyperlink")
    if url:
        link.set(qn("r:id"), paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True))
    if anchor:
        link.set(qn("w:anchor"), anchor)
    run = parse_xml(f"<w:r {nsdecls('w')}>"
                    f"{f'<w:rPr>{properties}</w:rPr>' if properties else ''}"
                    f'<w:t xml:space="preserve">{text}</w:t></w:r>')
    link.append(run)
    paragraph._p.append(link)


def _bookmark(paragraph, name: str, bookmark_id: int = 1) -> None:
    paragraph._p.insert(1, parse_xml(
        f'<w:bookmarkStart {nsdecls("w")} w:id="{bookmark_id}" w:name="{name}"/>'))
    paragraph._p.append(parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bookmark_id}"/>'))


def _add_numbering(document, abstract_xml: str, nums_xml: str) -> None:
    numbering = document.part.numbering_part.element
    abstract = parse_xml(abstract_xml)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:
        first_num.addprevious(abstract)
    else:
        numbering.append(abstract)
    for num in parse_xml(f"<w:x {nsdecls('w')}>{nums_xml}</w:x>"):
        numbering.append(num)


def _numbered(document, text: str, num_id: int, ilvl: int = 0, style: str | None = None):
    paragraph = document.add_paragraph(text, style=style)
    paragraph._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {nsdecls("w")}><w:ilvl w:val="{ilvl}"/>'
        f'<w:numId w:val="{num_id}"/></w:numPr>'))
    return paragraph


_MULTILEVEL = f"""
<w:abstractNum {nsdecls('w')} w:abstractNumId="90">
  <w:multiLevelType w:val="multilevel"/>
  <w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/>
    <w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>
  <w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="%2)"/>
    <w:pPr><w:ind w:left="1440" w:hanging="360"/></w:pPr></w:lvl>
  <w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="o"/>
    <w:pPr><w:ind w:left="2160" w:hanging="360"/></w:pPr></w:lvl>
</w:abstractNum>
"""


def _story_part(document, partname: str, content_type: str, reltype: str, xml_body: str,
                links: dict[str, str] | None = None) -> XmlPart:
    """A footnotes/endnotes part with its own hyperlink relationships."""
    part = XmlPart(PackURI(partname), content_type,
                   parse_xml(f"<w:x {nsdecls('w')}/>"), document.part.package)
    ids = {}
    for key, url in (links or {}).items():
        ids[key] = part.relate_to(url, RT.HYPERLINK, is_external=True)
    part._element = parse_xml(xml_body.format(**ids))
    document.part.relate_to(part, reltype)
    return part


def _caption(document, label: str, separator: str, text: str, sequence: str | None = None,
             simple: bool = False):
    """A numbered caption as Word and the forward converter write it."""
    paragraph = document.add_paragraph(style="Caption")
    if sequence is None:
        number = _xml_run("1")
    elif simple:
        number = f'<w:fldSimple w:instr=" SEQ {sequence} \\* ARABIC ">{_xml_run("1")}</w:fldSimple>'
    else:
        number = (
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            f'<w:r><w:instrText xml:space="preserve"> SEQ {sequence} \\* ARABIC </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>' + _xml_run("1")
            + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        )
    _append_xml(paragraph, _xml_run(f"{label} ") + number + _xml_run(separator) + _xml_run(text))
    return paragraph


def _quote_bar(paragraph, left_indent: int) -> None:
    """The forward converter's block-quote marking: a grey bar plus indent."""
    properties = paragraph._p.get_or_add_pPr()
    properties.append(parse_xml(
        f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="12" w:space="8" '
        f'w:color="A6A6A6"/></w:pBdr>'))
    properties.append(parse_xml(f'<w:ind {nsdecls("w")} w:left="{left_indent}"/>'))


class _ConverterTestCase(unittest.TestCase):
    def setUp(self) -> None:
        self._tmpdir = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmpdir.cleanup)
        self.root = Path(self._tmpdir.name)

    def convert(self, document, name: str = "doc", **options) -> tuple[str, list[str]]:
        source = self.root / f"{name}.docx"
        output = self.root / f"{name}.md"
        document.save(str(source))
        warnings = WordToMarkdownConverter(**options).convert_file(source, output)
        return output.read_text(encoding="utf-8"), warnings

    def codes(self, warnings: list[str]) -> list[str]:
        return [getattr(warning, "code", None) for warning in warnings]


class BlockOrderTests(_ConverterTestCase):
    def test_tables_stay_where_they_are_in_the_body(self) -> None:
        document = Document()
        document.add_paragraph("Before")
        table = document.add_table(rows=2, cols=2)
        for row, values in zip(table.rows, (("A", "B"), ("1", "2")), strict=True):
            for cell, value in zip(row.cells, values, strict=True):
                cell.text = value
        document.add_paragraph("After")

        markdown, warnings = self.convert(document)

        self.assertEqual(
            markdown,
            "Before\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n\nAfter\n",
        )
        self.assertEqual(warnings, [])

    def test_empty_paragraphs_collapse(self) -> None:
        document = Document()
        document.add_paragraph("One")
        document.add_paragraph("")
        document.add_paragraph("   ")
        document.add_paragraph("Two")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "One\n\nTwo\n")

    def test_content_controls_are_opened(self) -> None:
        document = Document()
        document.add_paragraph("Visible")
        sdt = parse_xml(
            f"<w:sdt {nsdecls('w')}><w:sdtPr/><w:sdtContent>"
            f"<w:p><w:r><w:t>Inside a control</w:t></w:r></w:p>"
            f"</w:sdtContent></w:sdt>"
        )
        document.element.body.insert(1, sdt)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "Visible\n\nInside a control\n")


class HeadingTests(_ConverterTestCase):
    def test_heading_levels_title_and_subtitle(self) -> None:
        document = Document()
        document.add_heading("Document title", level=0)
        document.add_paragraph("A subtitle", style="Subtitle")
        document.add_heading("Раздел", level=2)
        document.add_heading("Deep", level=6)

        markdown, warnings = self.convert(document)

        self.assertEqual(
            markdown,
            "# Document title\n\n*A subtitle*\n\n## Раздел\n\n###### Deep\n",
        )
        self.assertEqual(warnings, [])

    def test_heading_below_level_six_becomes_bold_with_a_warning(self) -> None:
        document = Document()
        document.add_paragraph("Seventh", style="Heading 7")
        document.add_paragraph("Eighth", style="Heading 8")

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "**Seventh**\n\n**Eighth**\n")
        self.assertEqual(self.codes(warnings), ["heading_level_clamped"])
        self.assertIn("(2)", warnings[0])

    def test_custom_style_with_outline_level_is_a_heading(self) -> None:
        document = Document()
        style = document.styles.add_style("Chapter", WD_STYLE_TYPE.PARAGRAPH)
        style.element.get_or_add_pPr().append(
            parse_xml(f'<w:outlineLvl {nsdecls("w")} w:val="2"/>'))
        document.add_paragraph("Custom", style="Chapter")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "### Custom\n")

    def test_uniform_bold_in_a_heading_is_style_not_emphasis(self) -> None:
        document = Document()
        heading = document.add_heading("", level=1)
        _run(heading, "All bold", bold=True)
        mixed = document.add_heading("", level=2)
        _run(mixed, "Some ")
        _run(mixed, "bold", bold=True)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "# All bold\n\n## Some **bold**\n")

    def test_trailing_hash_in_a_heading_is_escaped(self) -> None:
        document = Document()
        document.add_heading("Issue #", level=1)
        document.add_heading("C#", level=1)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "# Issue \\#\n\n# C#\n")


class InlineFormattingTests(_ConverterTestCase):
    def test_adjacent_runs_with_the_same_format_are_merged(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        _run(paragraph, "a", bold=True)
        _run(paragraph, "b", bold=True)
        _run(paragraph, " plain")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "**ab** plain\n")

    def test_whitespace_goes_outside_the_markers(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        _run(paragraph, "x")
        _run(paragraph, " bold ", bold=True)
        _run(paragraph, "y")
        _run(paragraph, "   ", italic=True)
        _run(paragraph, "z")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "x **bold** y   z\n")

    def test_nested_and_overlapping_emphasis(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        _run(paragraph, "both", bold=True, italic=True)
        _run(paragraph, " ")
        _run(paragraph, "bold ", bold=True)
        _run(paragraph, "and italic", bold=True, italic=True)
        _run(paragraph, " end", bold=True)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "***both*** **bold *and italic* end**\n")

    def test_strike_code_scripts_underline_and_highlight(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        _run(paragraph, "gone", strike=True)
        _run(paragraph, " ")
        _run(paragraph, "print(x)", font="Consolas")
        _run(paragraph, " E=mc")
        _run(paragraph, "2", superscript=True)
        _run(paragraph, " H")
        _run(paragraph, "2", subscript=True)
        _run(paragraph, "O ")
        _run(paragraph, "under", underline=True)
        _run(paragraph, " ")
        run = _run(paragraph, "marked")
        run._r.get_or_add_rPr().append(parse_xml(f'<w:highlight {nsdecls("w")} w:val="yellow"/>'))

        markdown, _ = self.convert(document)

        self.assertEqual(
            markdown,
            "~~gone~~ `print(x)` E=mc<sup>2</sup> H<sub>2</sub>O <u>under</u> ==marked==\n",
        )

    def test_inline_code_with_backticks_uses_a_longer_fence(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Use ")
        _run(paragraph, "a`b", font="Courier New")
        _run(paragraph, " or ")
        _run(paragraph, "`x`", font="Courier New")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "Use ``a`b`` or `` `x` ``\n")

    def test_code_character_style_counts_as_inline_code(self) -> None:
        document = Document()
        document.styles.add_style("Verbatim Char", WD_STYLE_TYPE.CHARACTER)
        paragraph = document.add_paragraph("Call ")
        paragraph.add_run("main()", style="Verbatim Char")
        paragraph.add_run(" now")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "Call `main()` now\n")

    def test_line_breaks_tabs_and_hidden_text(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        _run(paragraph, "first\tcolumn")
        paragraph.add_run().add_break()
        _run(paragraph, "# second line")
        hidden = _run(paragraph, " secret")
        hidden.font.hidden = True

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "first column\\\n\\# second line\n")

    def test_unparseable_emphasis_falls_back_to_html(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        _run(paragraph, "Note:", bold=True)
        _run(paragraph, "text and x")
        _run(paragraph, "(a)", italic=True)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "<strong>Note:</strong>text and x<em>(a)</em>\n")

    def test_all_caps_runs_are_upper_cased(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        _run(paragraph, "shout", all_caps=True)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "SHOUT\n")


class EscapingTests(_ConverterTestCase):
    SAMPLES = (
        "Costs $5 and $10",
        "*not bold* and _not italic_",
        "# not a heading",
        "1. not a list",
        "- not a bullet",
        "> not a quote",
        "a | b | c",
        "[not](a link) and ![no](image)",
        "back\\slash `tick` <tag> &amp; ~~strike~~ ==mark== x^2^",
        "snake_case_name stays readable",
        "===",
    )

    def test_special_characters_are_escaped(self) -> None:
        document = Document()
        for sample in self.SAMPLES[:3]:
            document.add_paragraph(sample)

        markdown, _ = self.convert(document)

        self.assertEqual(
            markdown,
            "Costs \\$5 and \\$10\n\n\\*not bold\\* and \\_not italic\\_\n\n\\# not a heading\n",
        )

    def test_escaped_text_parses_back_to_the_same_text(self) -> None:
        document = Document()
        for sample in self.SAMPLES:
            document.add_paragraph(sample)

        markdown, _ = self.convert(document)

        parser = MarkdownIt("commonmark").enable("table").enable("strikethrough")
        rendered = []
        for token in parser.parse(markdown):
            if token.type == "inline":
                rendered.append("".join(
                    child.content for child in token.children
                    if child.type in ("text", "code_inline")
                ))
        self.assertEqual(rendered, list(self.SAMPLES))
        self.assertIn("snake_case_name", markdown)


class HyperlinkTests(_ConverterTestCase):
    def test_external_links_keep_their_text(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("See ")
        _hyperlink(paragraph, "the docs", url="https://example.com/a b")
        paragraph.add_run(" and ")
        _hyperlink(paragraph, "https://example.com", url="https://example.com")
        paragraph.add_run(".")

        markdown, _ = self.convert(document)

        self.assertEqual(
            markdown,
            "See [the docs](<https://example.com/a b>) and <https://example.com>.\n",
        )

    def test_formatted_link_text(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        _hyperlink(paragraph, "bold link", url="https://example.com", properties="<w:b/>")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "**[bold link](https://example.com)**\n")

    def test_internal_link_to_a_heading_bookmark_uses_its_slug(self) -> None:
        document = Document()
        heading = document.add_heading("Введение и цели!", level=1)
        _bookmark(heading, "_Toc123")
        document.add_heading("Введение и цели!", level=1)
        paragraph = document.add_paragraph("Go to ")
        _hyperlink(paragraph, "intro", anchor="_Toc123")
        paragraph.add_run(" or ")
        _hyperlink(paragraph, "elsewhere", anchor="plain_mark")

        markdown, _ = self.convert(document)

        self.assertIn("[intro](#введение-и-цели)", markdown)
        self.assertIn("[elsewhere](#plain_mark)", markdown)

    def test_field_code_hyperlink(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Visit ")
        _append_xml(paragraph, (
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> HYPERLINK "https://example.org" </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>'
            + _xml_run("our site", '<w:rStyle w:val="Hyperlink"/><w:u w:val="single"/>')
            + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
            + _xml_run(" today.")
        ))

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "Visit [our site](https://example.org) today.\n")

    def test_simple_field_hyperlink(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        _append_xml(paragraph, (
            '<w:fldSimple w:instr=\' HYPERLINK "https://example.net" \'>'
            + _xml_run("net") + "</w:fldSimple>"
        ))

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "[net](https://example.net)\n")


class ListTests(_ConverterTestCase):
    def test_nested_numbered_list_with_counters_and_restart(self) -> None:
        document = Document()
        _add_numbering(document, _MULTILEVEL, (
            '<w:num w:numId="90"><w:abstractNumId w:val="90"/></w:num>'
            '<w:num w:numId="91"><w:abstractNumId w:val="90"/>'
            '<w:lvlOverride w:ilvl="0"><w:startOverride w:val="5"/></w:lvlOverride></w:num>'
        ))
        _numbered(document, "first", 90)
        _numbered(document, "nested a", 90, 1)
        _numbered(document, "nested b", 90, 1)
        _numbered(document, "deep bullet", 90, 2)
        _numbered(document, "second", 90)
        _numbered(document, "nested again", 90, 1)
        document.add_paragraph("Between lists.")
        _numbered(document, "restarted", 91)
        _numbered(document, "continues", 91)

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, (
            "1. first\n"
            "   1. nested a\n"
            "   2. nested b\n"
            "      - deep bullet\n"
            "2. second\n"
            "   1. nested again\n"
            "\n"
            "Between lists.\n"
            "\n"
            "5. restarted\n"
            "6. continues\n"
        ))
        self.assertEqual(warnings, [])

    def test_style_based_lists(self) -> None:
        document = Document()
        document.add_paragraph("one", style="List Bullet")
        document.add_paragraph("child", style="List Bullet 2")
        document.add_paragraph("two", style="List Bullet")
        document.add_paragraph("Text.")
        document.add_paragraph("alpha", style="List Number")
        document.add_paragraph("beta", style="List Number")

        markdown, _ = self.convert(document)

        self.assertEqual(
            markdown,
            "- one\n  - child\n- two\n\nText.\n\n1. alpha\n2. beta\n",
        )

    def test_task_list_glyphs(self) -> None:
        document = Document()
        document.add_paragraph("☐ todo", style="List Bullet")
        document.add_paragraph("☒ done", style="List Bullet")
        document.add_paragraph("☑done too", style="List Bullet")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "- [ ] todo\n- [x] done\n- [x] done too\n")

    def test_adjacent_separate_lists_get_distinct_markers(self) -> None:
        document = Document()
        _add_numbering(document, _MULTILEVEL.replace('"90"', '"92"'), (
            '<w:num w:numId="92"><w:abstractNumId w:val="92"/></w:num>'
            '<w:num w:numId="93"><w:abstractNumId w:val="92"/></w:num>'
        ))
        _numbered(document, "a", 92)
        _numbered(document, "b", 92)
        _numbered(document, "c", 93)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "1. a\n2. b\n\n1) c\n")
        parsed = MarkdownIt("commonmark").parse(markdown)
        self.assertEqual(sum(t.type == "ordered_list_open" for t in parsed), 2)

    def test_indented_paragraph_continues_the_list_item(self) -> None:
        document = Document()
        document.add_paragraph("item", style="List Bullet")
        document.add_paragraph("more about the item", style="List Continue")
        document.add_paragraph("next", style="List Bullet")
        document.add_paragraph("Not indented.")

        markdown, _ = self.convert(document)

        self.assertEqual(
            markdown,
            "- item\n\n  more about the item\n\n- next\n\nNot indented.\n",
        )

    def test_indented_code_block_and_table_stay_inside_the_item(self) -> None:
        document = Document()
        document.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
        first = document.add_paragraph("Install", style="List Number")
        code = document.add_paragraph("pip install x", style="Source Code")
        code.paragraph_format.left_indent = Pt(18)  # the item's text indent: 360 twips
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = "cell"
        table._tbl.tblPr.append(parse_xml(
            f'<w:tblInd {nsdecls("w")} w:w="360" w:type="dxa"/>'))
        document.add_paragraph("Run", style="List Number")
        self.assertIsNotNone(first)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, (
            "1. Install\n\n"
            "   ```\n   pip install x\n   ```\n\n"
            "   | cell |\n   | --- |\n\n"
            "2. Run\n"
        ))

    def test_numbered_headings_keep_their_number(self) -> None:
        document = Document()
        # Level 2 counts in letters but is "legal": printed as "1.1", not "1.a".
        _add_numbering(document, _MULTILEVEL.replace('"90"', '"94"').replace(
            '<w:lvlText w:val="%2)"/>', '<w:isLgl/><w:lvlText w:val="%1.%2"/>'), (
            '<w:num w:numId="94"><w:abstractNumId w:val="94"/></w:num>'))
        _numbered(document, "Intro", 94, 0, style="Heading 1")
        _numbered(document, "Scope", 94, 1, style="Heading 2")
        _numbered(document, "Terms", 94, 1, style="Heading 2")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "# 1. Intro\n\n## 1.1 Scope\n\n## 1.2 Terms\n")

    def test_links_to_numbered_headings_use_the_numbered_anchor(self) -> None:
        document = Document()
        _add_numbering(document, _MULTILEVEL.replace('"90"', '"95"'), (
            '<w:num w:numId="95"><w:abstractNumId w:val="95"/></w:num>'))
        paragraph = document.add_paragraph("See ")
        _hyperlink(paragraph, "methods", anchor="_Toc2")
        _numbered(document, "Intro", 95, 0, style="Heading 1")
        methods = _numbered(document, "Methods", 95, 0, style="Heading 1")
        _bookmark(methods, "_Toc2")

        markdown, _ = self.convert(document)

        self.assertIn("[methods](#2-methods)", markdown)
        self.assertIn("# 2. Methods", markdown)

    def test_hard_break_inside_a_list_item_keeps_its_indent(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("line one", style="List Number")
        paragraph.add_run().add_break()
        paragraph.add_run("1. line two")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "1. line one\\\n   1\\. line two\n")


class CodeBlockTests(_ConverterTestCase):
    def test_monospace_paragraphs_become_one_fenced_block(self) -> None:
        document = Document()
        document.add_paragraph("Intro")
        for line in ("def f():", "\treturn `x`"):
            paragraph = document.add_paragraph()
            _run(paragraph, line, font="Courier New")
        document.add_paragraph("Outro")

        markdown, _ = self.convert(document)

        self.assertEqual(
            markdown,
            "Intro\n\n```\ndef f():\n\treturn `x`\n```\n\nOutro\n",
        )

    def test_source_code_style_keeps_blank_lines_and_indentation(self) -> None:
        document = Document()
        document.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
        for line in ("if x:", "    y = 1", "", "print(y)"):
            document.add_paragraph(line, style="Source Code")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "```\nif x:\n    y = 1\n\nprint(y)\n```\n")

    def test_legacy_caption_becomes_the_info_string(self) -> None:
        document = Document()
        caption = document.add_paragraph()
        _run(caption, "python", italic=True)
        code = document.add_paragraph()
        run = _run(code, "a = 1\nb = '``` fence'", font="Courier New")
        run.font.size = Pt(10)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "````python\na = 1\nb = '``` fence'\n````\n")

    def test_italic_word_without_code_after_it_stays_text(self) -> None:
        document = Document()
        caption = document.add_paragraph()
        _run(caption, "python", italic=True)
        document.add_paragraph("Ordinary text.")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "*python*\n\nOrdinary text.\n")


class QuoteAndRuleTests(_ConverterTestCase):
    def test_consecutive_quote_paragraphs_form_one_block_quote(self) -> None:
        document = Document()
        document.add_paragraph("First quoted.", style="Quote")
        document.add_paragraph("Second quoted.", style="Intense Quote")
        document.styles.add_style("Quote 2", WD_STYLE_TYPE.PARAGRAPH)
        document.add_paragraph("Nested.", style="Quote 2")
        document.add_paragraph("After.")

        markdown, _ = self.convert(document)

        self.assertEqual(
            markdown,
            "> First quoted.\n>\n> Second quoted.\n>\n> > Nested.\n\nAfter.\n",
        )

    def test_shaded_call_out_with_an_alert_title_becomes_a_github_alert(self) -> None:
        document = Document()
        for text, bold in (("Примечание", True), ("Важное замечание.", False)):
            paragraph = document.add_paragraph()
            _run(paragraph, text, bold=bold)
            paragraph._p.get_or_add_pPr().append(parse_xml(
                f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="4" '
                f'w:color="0969DA"/></w:pBdr>'))
            paragraph._p.get_or_add_pPr().append(parse_xml(
                f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="DDF4FF"/>'))
        document.add_paragraph("After.")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "> [!NOTE]\n> Важное замечание.\n\nAfter.\n")

    def test_quote_nesting_follows_direct_indentation(self) -> None:
        document = Document()
        for text, indent in (("outer", 18), ("inner", 36), ("outer again", 18)):
            paragraph = document.add_paragraph(text, style="Quote")
            paragraph.paragraph_format.left_indent = Pt(indent)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "> outer\n>\n> > inner\n>\n> outer again\n")

    def test_lists_and_code_inside_a_block_quote_keep_the_quote(self) -> None:
        document = Document()
        document.styles.add_style("Source Code", WD_STYLE_TYPE.PARAGRAPH)
        _add_numbering(document, _MULTILEVEL.replace('"90"', '"96"'), (
            '<w:num w:numId="96"><w:abstractNumId w:val="96"/></w:num>'
            '<w:num w:numId="97"><w:abstractNumId w:val="96"/></w:num>'))
        _quote_bar(document.add_paragraph("Quoted text", style="Quote"), 360)
        for text in ("quoted item", "second"):
            # Level 0 hangs at 720; one quote level adds 360.
            _quote_bar(_numbered(document, text, 96, 2, style="List Paragraph"), 2160 + 360)
        _quote_bar(_numbered(document, "deeper numbered", 97, 0, style="List Paragraph"),
                   720 + 2 * 360)
        _quote_bar(document.add_paragraph("quoted code", style="Source Code"), 360)
        document.add_paragraph("Plain after.")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, (
            "> Quoted text\n>\n"
            "> - quoted item\n> - second\n>\n"
            "> > 1. deeper numbered\n>\n"
            "> ```\n> quoted code\n> ```\n\n"
            "Plain after.\n"
        ))
        parsed = MarkdownIt("commonmark").parse(markdown)
        self.assertEqual([t.type for t in parsed if t.type.endswith("_open")][:3],
                         ["blockquote_open", "paragraph_open", "bullet_list_open"])

    def test_bordered_empty_paragraph_is_a_thematic_break(self) -> None:
        document = Document()
        document.add_paragraph("Above")
        rule = document.add_paragraph()
        rule._p.get_or_add_pPr().append(parse_xml(
            f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" '
            f'w:color="808080"/></w:pBdr>'))
        document.add_paragraph("Below")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "Above\n\n---\n\nBelow\n")


class TableTests(_ConverterTestCase):
    def test_cells_keep_formatting_breaks_and_escape_pipes(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=3)
        header = table.rows[0].cells
        header[0].text = "Left"
        header[1].text = "Center"
        header[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        header[2].text = "Right"
        header[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        body = table.rows[1].cells
        _run(body[0].paragraphs[0], "bold", bold=True)
        body[0].paragraphs[0].add_run().add_break()
        body[0].paragraphs[0].add_run("x | y")
        _run(body[1].paragraphs[0], "a|b", font="Courier New")
        body[2].text = "one"
        body[2].add_paragraph("two")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, (
            "| Left | Center | Right |\n"
            "| --- | :---: | ---: |\n"
            "| **bold**<br>x \\| y | `a\\|b` | one<br>two |\n"
        ))
        cells = MarkdownIt("commonmark").enable("table").render(markdown)
        self.assertIn("<code>a|b</code>", cells)
        self.assertIn("x | y", cells)

    def test_bold_header_cells_lose_the_redundant_markers(self) -> None:
        document = Document()
        table = document.add_table(rows=2, cols=2)
        _run(table.cell(0, 0).paragraphs[0], "Name", bold=True)
        cell = table.cell(0, 1).paragraphs[0]
        _run(cell, "Mixed ")
        _run(cell, "bold", bold=True)
        _run(table.cell(1, 0).paragraphs[0], "body", bold=True)
        table.cell(1, 1).text = "x"

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, (
            "| Name | Mixed **bold** |\n| --- | --- |\n| **body** | x |\n"
        ))

    def test_merged_cells_repeat_empty_cells_and_warn_once(self) -> None:
        document = Document()
        for _ in range(2):
            table = document.add_table(rows=2, cols=3)
            table.cell(0, 0).merge(table.cell(0, 1)).text = "wide"
            table.cell(0, 2).text = "c"
            table.cell(1, 0).text = "1"
            table.cell(1, 1).text = "2"
            table.cell(1, 2).text = "3"

        markdown, warnings = self.convert(document)

        self.assertIn("| wide |  | c |\n| --- | --- | --- |\n| 1 | 2 | 3 |", markdown)
        self.assertEqual(self.codes(warnings), ["table_merged_cells"])
        self.assertIn("2 table(s)", warnings[0])

    def test_single_row_table_still_gets_a_separator(self) -> None:
        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "only"
        table.cell(0, 1).text = "row"

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "| only | row |\n| --- | --- |\n")

    def test_nested_table_is_flattened_with_a_warning(self) -> None:
        document = Document()
        table = document.add_table(rows=1, cols=1)
        inner = table.cell(0, 0).add_table(rows=1, cols=2)
        inner.cell(0, 0).text = "in1"
        inner.cell(0, 1).text = "in2"

        markdown, warnings = self.convert(document)

        self.assertIn("in1; in2", markdown)
        self.assertEqual(self.codes(warnings), ["nested_table_flattened"])


class ImageTests(_ConverterTestCase):
    def test_pictures_are_extracted_next_to_the_output(self) -> None:
        document = Document()
        document.add_paragraph("Red:")
        shape = document.add_picture(BytesIO(_png((255, 0, 0))))
        shape._inline.docPr.set("descr", "A red [dot]")
        document.add_picture(BytesIO(_png((0, 0, 255))))
        document.add_picture(BytesIO(_png((255, 0, 0))))  # same image again

        markdown, warnings = self.convert(document, name="report")

        self.assertEqual(markdown, (
            "Red:\n\n"
            "![A red \\[dot\\]](report_media/image1.png)\n\n"
            "![](report_media/image2.png)\n\n"
            "![](report_media/image1.png)\n"
        ))
        media = self.root / "report_media"
        self.assertEqual(sorted(path.name for path in media.iterdir()),
                         ["image1.png", "image2.png"])
        self.assertEqual((media / "image1.png").read_bytes(), _png((255, 0, 0)))
        self.assertEqual(warnings, [])

    def test_no_media_folder_without_images(self) -> None:
        document = Document()
        document.add_paragraph("Text only")

        self.convert(document, name="plain")

        self.assertFalse((self.root / "plain_media").exists())

    def test_extract_media_false_keeps_alt_text_and_warns(self) -> None:
        document = Document()
        shape = document.add_picture(BytesIO(_png()))
        shape._inline.docPr.set("descr", "chart")

        markdown, warnings = self.convert(document, name="noimg", extract_media=False)

        self.assertEqual(markdown, "chart\n")
        self.assertEqual(self.codes(warnings), ["image_not_extracted"])
        self.assertFalse((self.root / "noimg_media").exists())

    def test_svg_is_preferred_over_its_png_fallback(self) -> None:
        document = Document()
        shape = document.add_picture(BytesIO(_png()))
        svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="1" height="1"/>'
        from docx.opc.part import Part
        svg_part = Part(PackURI("/word/media/vector1.svg"), "image/svg+xml", svg,
                        document.part.package)
        relationship_id = document.part.relate_to(svg_part, RT.IMAGE)
        blip = shape._inline.find(".//" + qn("a:blip"))
        blip.append(parse_xml(
            f'<a:extLst {nsdecls("a")}><a:ext uri="{{96DAC541-7B7A-43D3-8B79-37D633B846F1}}">'
            f'<asvg:svgBlip xmlns:asvg="http://schemas.microsoft.com/office/drawing/2016/SVG/main" '
            f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
            f'r:embed="{relationship_id}"/></a:ext></a:extLst>'))

        markdown, _ = self.convert(document, name="vec")

        self.assertEqual(markdown, "![](vec_media/image1.svg)\n")
        self.assertEqual((self.root / "vec_media" / "image1.svg").read_bytes(), svg)

    def test_linked_picture_keeps_its_url(self) -> None:
        document = Document()
        shape = document.add_picture(BytesIO(_png()))
        blip = shape._inline.find(".//" + qn("a:blip"))
        relationship_id = document.part.relate_to(
            "https://example.com/pic.png", RT.IMAGE, is_external=True)
        del blip.attrib[qn("r:embed")]
        blip.set(qn("r:link"), relationship_id)

        markdown, _ = self.convert(document, name="linked")

        self.assertEqual(markdown, "![](https://example.com/pic.png)\n")

    def test_text_boxes_and_vml_shapes_are_reported(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Shape:")
        _append_xml(paragraph, (
            f'<w:r><w:drawing><wp:anchor xmlns:wp="http://schemas.openxmlformats.org/'
            f'drawingml/2006/wordprocessingDrawing"><a:graphic xmlns:a="http://schemas.'
            f'openxmlformats.org/drawingml/2006/main"><a:graphicData uri="{_WPS}">'
            f'<wps:wsp xmlns:wps="{_WPS}"><wps:txbx><w:txbxContent><w:p><w:r><w:t>boxed'
            f'</w:t></w:r></w:p></w:txbxContent></wps:txbx></wps:wsp></a:graphicData>'
            f'</a:graphic></wp:anchor></w:drawing></w:r>'
            f'<w:r><w:pict><v:oval xmlns:v="urn:schemas-microsoft-com:vml"/></w:pict></w:r>'
        ))

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "Shape:\n")
        self.assertEqual(sorted(self.codes(warnings)), ["image_unsupported", "textbox_skipped"])


class CaptionTests(_ConverterTestCase):
    def _picture(self, document, alt: str = "") -> None:
        shape = document.add_picture(BytesIO(_png()))
        if alt:
            shape._inline.docPr.set("descr", alt)
        document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def test_figure_caption_becomes_the_image_title(self) -> None:
        document = Document()
        self._picture(document, alt="Схема")
        _caption(document, "Рисунок", " — ", "Подпись картинки", sequence="Figure")
        document.add_paragraph("After.")

        markdown, warnings = self.convert(document, name="fig")

        self.assertEqual(
            markdown, '![Схема](fig_media/image1.png "Подпись картинки")\n\nAfter.\n')
        self.assertEqual(warnings, [])

    def test_caption_equal_to_alt_or_without_alt_needs_no_title(self) -> None:
        document = Document()
        self._picture(document, alt="SVG")
        _caption(document, "Figure", ": ", "SVG", sequence="Figure")
        self._picture(document)
        _caption(document, "Figure", ": ", 'A "quoted" caption', sequence="Figure", simple=True)

        markdown, _ = self.convert(document, name="fig")

        self.assertEqual(markdown, (
            "![SVG](fig_media/image1.png)\n\n"
            '![A "quoted" caption](fig_media/image1.png)\n'
        ))

    def test_table_caption_becomes_a_table_line(self) -> None:
        document = Document()
        _caption(document, "Таблица", " — ", "Итоги года", sequence="Table")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "x"
        document.add_table(rows=1, cols=1).cell(0, 0).text = "y"
        _caption(document, "Table", ". ", "Below *it*", simple=True, sequence="Table")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, (
            "Таблица: Итоги года\n\n| x |\n| --- |\n\n"
            "| y |\n| --- |\n\nTable: Below \\*it\\*\n"
        ))

    def test_other_caption_paragraphs_stay_paragraphs(self) -> None:
        document = Document()
        self._picture(document)
        document.add_paragraph("Just a styled note", style="Caption")
        _caption(document, "Figure", ": ", "Far from any picture", sequence="Figure")

        markdown, _ = self.convert(document, name="cap")

        self.assertEqual(markdown, (
            "![](cap_media/image1.png)\n\nJust a styled note\n\nFigure 1: Far from any picture\n"
        ))

    def test_typed_caption_without_a_field_needs_a_known_label(self) -> None:
        document = Document()
        self._picture(document)
        _caption(document, "Рис.", " – ", "Без поля")
        self._picture(document)
        _caption(document, "Step", ": ", "Not a caption label")

        markdown, _ = self.convert(document, name="typed")

        self.assertIn('![Без поля](typed_media/image1.png)', markdown)
        self.assertIn("Step 1: Not a caption label", markdown)


class EquationTests(_ConverterTestCase):
    def test_inline_equation_becomes_dollar_math(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Energy ")
        paragraph._p.append(latex_to_omml(r"E = mc^2"))
        paragraph.add_run(" holds.")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "Energy $E = mc^{2}$ holds.\n")

    def test_display_equation_with_label(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph._p.append(latex_to_omml(r"\frac{a}{b}"))
        paragraph.add_run("\t(1)")
        para = document.add_paragraph()
        math_paragraph = parse_xml(f"<m:oMathPara {nsdecls('m')}/>")
        math_paragraph.append(latex_to_omml(r"\sqrt{x}"))
        para._p.append(math_paragraph)

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "$$\n\\frac{a}{b}\n$$ (1)\n\n$$\n\\sqrt{x}\n$$\n")

    def test_multi_line_equation_keeps_its_alignment(self) -> None:
        document = Document()
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph._p.append(latex_to_omml(r"a &= b + c \\ d &= e"))

        markdown, _ = self.convert(document)

        self.assertEqual(
            markdown,
            "$$\n\\begin{aligned}\na &= b+c \\\\\nd &= e\n\\end{aligned}\n$$\n",
        )

    def test_unknown_equation_element_keeps_text_and_warns(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Odd ")
        paragraph._p.append(parse_xml(
            f"<m:oMath {nsdecls('m')}><m:weird><m:r><m:t>q</m:t></m:r></m:weird></m:oMath>"))

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "Odd $q$\n")
        self.assertEqual(self.codes(warnings), ["formula_partial"])
        self.assertIn("m:weird", warnings[0])


class NoteTests(_ConverterTestCase):
    FOOTNOTES = (
        f"<w:footnotes {nsdecls('w', 'r')}>"
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r>'
        "<w:continuationSeparator/></w:r></w:p></w:footnote>"
        '<w:footnote w:id="1"><w:p><w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
        "<w:footnoteRef/></w:r>"
        + _xml_run(" See ") + _xml_run("bold", "<w:b/>") + _xml_run(" and ")
        + '<w:hyperlink r:id="{docs}">' + _xml_run("the docs") + "</w:hyperlink>"
        + _xml_run(".") + "</w:p><w:p>" + _xml_run("Second paragraph.")
        + "</w:p></w:footnote>"
        '<w:footnote w:id="2"><w:p>' + _xml_run("Short note") + "</w:p></w:footnote>"
        "</w:footnotes>"
    )

    def _reference(self, paragraph, note_id: int, kind: str = "footnote") -> None:
        _append_xml(paragraph, (
            '<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr>'
            f'<w:{kind}Reference w:id="{note_id}"/></w:r>'
        ))

    def test_footnotes_render_with_formatting_and_their_own_links(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Claim")
        self._reference(paragraph, 2)
        paragraph.add_run(" and another")
        self._reference(paragraph, 1)
        paragraph.add_run(" (aside).")
        _story_part(document, "/word/footnotes.xml", CT.WML_FOOTNOTES, RT.FOOTNOTES,
                    self.FOOTNOTES, {"docs": "https://example.com/docs"})

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, (
            "Claim[^2] and another[^1] (aside).\n\n"
            "[^2]: Short note\n\n"
            "[^1]: See **bold** and [the docs](https://example.com/docs).\n"
            "\n"
            "    Second paragraph.\n"
        ))
        self.assertEqual(warnings, [])

    def test_endnotes_get_their_own_labels(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("End")
        self._reference(paragraph, 1, kind="endnote")
        _story_part(
            document, "/word/endnotes.xml", CT.WML_ENDNOTES, RT.ENDNOTES,
            f"<w:endnotes {nsdecls('w')}>"
            '<w:endnote w:type="separator" w:id="0"><w:p/></w:endnote>'
            '<w:endnote w:id="1"><w:p>' + _xml_run("An endnote.") + "</w:p></w:endnote>"
            "</w:endnotes>",
        )

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "End[^e1]\n\n[^e1]: An endnote.\n")

    def test_reference_without_a_note_is_reported(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Dangling")
        self._reference(paragraph, 7)

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "Dangling[^7]\n")
        self.assertEqual(self.codes(warnings), ["footnote_missing"])


class RevisionAndFieldTests(_ConverterTestCase):
    def test_tracked_insertions_count_and_deletions_do_not(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Keep ")
        _append_xml(paragraph, (
            '<w:ins w:id="1" w:author="a">' + _xml_run("new") + "</w:ins>"
            '<w:del w:id="2" w:author="a"><w:r><w:delText>old</w:delText></w:r></w:del>'
            + _xml_run(" text")
        ))

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, "Keep new text\n")

    def test_table_of_contents_field_becomes_a_toc_marker(self) -> None:
        document = Document()
        first = document.add_paragraph()
        _append_xml(first, (
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" \\h </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>' + _xml_run("Intro 1")
        ))
        document.add_paragraph("Chapter 2", style="List Bullet")
        last = document.add_paragraph()
        _append_xml(last, '<w:r><w:fldChar w:fldCharType="end"/></w:r>')
        document.add_heading("Intro", level=1)
        page = document.add_paragraph("Page ")
        _append_xml(page, '<w:fldSimple w:instr=" PAGE ">' + _xml_run("4") + "</w:fldSimple>")

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "[TOC]\n\n# Intro\n\nPage\n")
        self.assertEqual(warnings, [])

    def test_contents_title_before_a_table_of_contents_goes_too(self) -> None:
        document = Document()
        title = document.add_paragraph()
        _run(title, "Содержание", bold=True)
        toc = document.add_paragraph()
        _append_xml(toc, (
            '<w:r><w:fldChar w:fldCharType="begin"/></w:r>'
            '<w:r><w:instrText xml:space="preserve"> TOC \\o "1-3" </w:instrText></w:r>'
            '<w:r><w:fldChar w:fldCharType="separate"/></w:r>' + _xml_run("Введение")
            + '<w:r><w:fldChar w:fldCharType="end"/></w:r>'
        ))
        document.add_heading("Введение", level=1)

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "[TOC]\n\n# Введение\n")
        self.assertEqual(warnings, [])

    def test_table_of_contents_control_becomes_a_toc_marker(self) -> None:
        document = Document()
        document.add_paragraph("Before")
        document.element.body.insert(1, parse_xml(
            f"<w:sdt {nsdecls('w')}><w:sdtPr><w:docPartObj>"
            '<w:docPartGallery w:val="Table of Contents"/><w:docPartUnique/>'
            "</w:docPartObj></w:sdtPr><w:sdtContent>"
            '<w:p><w:pPr><w:pStyle w:val="TOCHeading"/></w:pPr>' + _xml_run("Contents") + "</w:p>"
            "<w:p>" + _xml_run("Intro 1") + "</w:p></w:sdtContent></w:sdt>"
        ))
        document.add_heading("Intro", level=1)

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "Before\n\n[TOC]\n\n# Intro\n")
        self.assertEqual(warnings, [])

    def test_table_of_figures_and_a_toc_inside_a_cell_are_reported(self) -> None:
        document = Document()
        figures = document.add_paragraph()
        _append_xml(figures, (
            '<w:fldSimple w:instr=\' TOC \\h \\c "Figure" \'>' + _xml_run("Figure 1 ... 3")
            + "</w:fldSimple>"
        ))
        cell = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
        _append_xml(cell, '<w:fldSimple w:instr=" TOC \\o ">' + _xml_run("Intro") + "</w:fldSimple>")

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "|  |\n| --- |\n")
        self.assertEqual(self.codes(warnings), ["toc_skipped"])
        self.assertIn("(2)", warnings[0])

    def test_table_of_contents_control_inside_a_cell_is_reported(self) -> None:
        def toc_control():
            return parse_xml(
                f"<w:sdt {nsdecls('w')}><w:sdtPr><w:docPartObj>"
                '<w:docPartGallery w:val="Table of Contents"/><w:docPartUnique/>'
                "</w:docPartObj></w:sdtPr><w:sdtContent>"
                "<w:p>" + _xml_run("Intro 1") + "</w:p></w:sdtContent></w:sdt>"
            )

        document = Document()
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).paragraphs[0].add_run("Before")
        table.cell(0, 0)._tc.append(toc_control())
        nested = table.cell(0, 1).add_table(rows=1, cols=1).cell(0, 0)
        nested.paragraphs[0].add_run("Inner")
        nested._tc.append(toc_control())

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "| Before | Inner |\n| --- | --- |\n")
        self.assertEqual(self.codes(warnings), ["toc_skipped", "nested_table_flattened"])
        self.assertIn("(2)", warnings[0])

    def test_symbol_characters(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Angle ")
        _append_xml(paragraph, '<w:r><w:sym w:font="Symbol" w:char="F061"/></w:r>'
                               '<w:r><w:sym w:font="Wingdings" w:char="F0FC"/></w:r>')

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "Angle α✓\n")
        self.assertEqual(warnings, [])

    def test_symbol_from_an_unknown_font_is_reported(self) -> None:
        document = Document()
        paragraph = document.add_paragraph("Odd ")
        _append_xml(paragraph, '<w:r><w:sym w:font="Webdings" w:char="F021"/></w:r>')

        markdown, warnings = self.convert(document)

        self.assertEqual(markdown, "Odd\n")
        self.assertEqual(self.codes(warnings), ["symbol_unsupported"])


class MetadataTests(_ConverterTestCase):


    def test_core_properties_become_front_matter(self) -> None:
        document = Document()
        document.core_properties.title = 'A "quoted" title'
        document.core_properties.author = "Иван Петров"
        document.core_properties.keywords = "docx, markdown"
        document.add_paragraph("Body")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, (
            '---\ntitle: "A \\"quoted\\" title"\nauthor: "Иван Петров"\n'
            'keywords: "docx, markdown"\n---\n\nBody\n'
        ))

    def test_title_page_made_from_metadata_is_not_repeated(self) -> None:
        document = Document()
        document.core_properties.title = "Отчёт"
        document.core_properties.author = "Иван"
        document.add_paragraph("Отчёт", style="Title")
        author = document.add_paragraph("Иван")
        author.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.add_paragraph("Body")

        markdown, _ = self.convert(document)

        self.assertEqual(markdown, '---\ntitle: "Отчёт"\nauthor: "Иван"\n---\n\nBody\n')

    def test_document_language_goes_into_front_matter(self) -> None:
        document = Document()
        document.core_properties.title = "T"
        document.core_properties.language = "ru-RU"
        document.add_paragraph("Body")
        alone = Document()
        alone.core_properties.language = "en-US"
        alone.add_paragraph("Body")

        markdown, _ = self.convert(document)
        lone, _ = self.convert(alone, name="alone")

        self.assertEqual(markdown, '---\ntitle: "T"\nlang: "ru-RU"\n---\n\nBody\n')
        self.assertEqual(lone, '---\nlang: "en-US"\n---\n\nBody\n')

    def test_front_matter_can_be_disabled(self) -> None:
        document = Document()
        document.core_properties.title = "T"
        document.add_paragraph("Body")

        markdown, _ = self.convert(document, front_matter=False)

        self.assertEqual(markdown, "Body\n")


class ContractTests(_ConverterTestCase):
    def test_missing_input_raises_conversion_error(self) -> None:
        with self.assertRaises(ConversionError):
            WordToMarkdownConverter().convert_file(self.root / "missing.docx",
                                                   self.root / "out.md")

    def test_not_a_docx_raises_conversion_error(self) -> None:
        source = self.root / "fake.docx"
        source.write_bytes(b"not a zip")
        with self.assertRaises(ConversionError):
            WordToMarkdownConverter().convert_file(source, self.root / "out.md")

    def test_warnings_are_conversion_warning_instances(self) -> None:
        document = Document()
        document.add_paragraph("Deep", style="Heading 9")

        _, warnings = self.convert(document)

        self.assertTrue(all(isinstance(w, ConversionWarning) for w in warnings))

    def test_convert_to_string_without_media_dir(self) -> None:
        document = Document()
        document.add_heading("Title", level=1)
        source = self.root / "s.docx"
        document.save(str(source))

        markdown, warnings = WordToMarkdownConverter().convert_to_string(source)

        self.assertEqual(markdown, "# Title\n")
        self.assertEqual(warnings, [])


class EndToEndTests(_ConverterTestCase):
    SAMPLE = (
        "# Main title\n\n"
        "Some **bold** and *italic* text with a [link](https://example.com).\n\n"
        "## Data\n\n"
        "| Name | Value |\n|---|---|\n| alpha | 1 |\n| beta | 2 |\n\n"
        "Inline math $x^2 + y^2 = z^2$ in a sentence.\n\n"
        "$$\n\\frac{a}{b}\n$$\n\n"
        "- first item\n- second item\n"
    )

    def test_markdown_survives_a_round_trip_through_word(self) -> None:
        from mdtoword.converters import MarkdownToWordConverter

        docx_path = self.root / "round.docx"
        MarkdownToWordConverter().convert_content(self.SAMPLE, docx_path)
        markdown_path = self.root / "round.md"
        WordToMarkdownConverter().convert_file(docx_path, markdown_path)
        markdown = markdown_path.read_text(encoding="utf-8")

        for expected in ("# Main title", "## Data", "**bold**", "*italic*",
                         "[link](https://example.com)", "| alpha | 1 |", "| beta | 2 |",
                         "$x^{2}", "\\frac{a}{b}", "first item", "second item"):
            self.assertIn(expected, markdown)
        self.assertRegex(markdown, r"(?m)^[-*] first item$")
        self.assertRegex(markdown, r"(?m)^\$\$$")

        # And the result converts back without losing that content again.
        again = self.root / "again.docx"
        MarkdownToWordConverter().convert_content(markdown, again)
        text = "\n".join(p.text for p in Document(str(again)).paragraphs)
        for expected in ("Main title", "bold", "italic", "link", "first item"):
            self.assertIn(expected, text)
        self.assertEqual(len(Document(str(again)).tables), 1)

    def test_captions_and_quoted_lists_do_not_change_on_a_second_round_trip(self) -> None:
        from mdtoword.converters import MarkdownToWordConverter

        (self.root / "dot.png").write_bytes(_png())
        source = (
            "Table: Итоги\n\n| A | B |\n|---|---|\n| 1 | 2 |\n\n"
            '![Точка](dot.png "Подпись")\n\n'
            "> Quote\n>\n> - one\n> - two\n>\n> > 1. deep\n\nEnd.\n"
        )
        (self.root / "r0.md").write_text(source, encoding="utf-8")
        texts = []
        for step in range(2):
            MarkdownToWordConverter().convert_file(self.root / f"r{step}.md",
                                                   self.root / f"r{step}.docx")
            WordToMarkdownConverter().convert_file(self.root / f"r{step}.docx",
                                                   self.root / f"r{step + 1}.md")
            text = (self.root / f"r{step + 1}.md").read_text(encoding="utf-8")
            texts.append(text.replace(f"r{step + 1}_media", "media"))
        self.assertEqual(texts[0], texts[1])
        for expected in ("Таблица: Итоги", '"Подпись")', "> - one", "> > 1. deep"):
            self.assertIn(expected, texts[0])


if __name__ == "__main__":
    unittest.main()
