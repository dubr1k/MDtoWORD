"""Math: inline and display formulas, amsmath environments, numbering.

Formulas become native Word equations through ``latex_omml``. Display
equations may carry a number (``$$ … $$ (1)`` or ``\\tag``), laid out with
the formula centred and the number flush right. Inline fragments that read
as prose rather than math, and formulas ``latex_omml`` cannot convert, are
kept verbatim with a warning -- a formula never silently vanishes.
"""

from __future__ import annotations

from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt

from ..latex_omml import UnsupportedLatexError, latex_to_omml
from .constants import (
    _ALIGNMENT_MARKER,
    _AMSMATH_WRAPPER,
    _CODE_FONT,
    _COLUMN_ARGUMENT,
    _COLUMN_ARGUMENT_ENVIRONMENTS,
    _CYRILLIC,
    _LINE_BREAK,
    _MATH_INDICATOR,
    _MATRIX_ENVIRONMENTS,
    _TEXT_COMMAND,
)
from .helpers import _equation_number, _plain_formatting, _split_tag
from .state import RendererState


class MathMixin(RendererState):
    """Render formulas as Word equations, or verbatim with a warning."""

    def _render_math(
        self, latex: str, display: bool, markup: str = "$", label: str | None = None
    ) -> None:
        """Insert a real Word equation, falling back to verbatim text.

        Every path ends in either an equation or the untouched source plus a
        warning, so a formula can never silently vanish or be silently wrong.
        """
        formula = latex.strip("\n").strip()
        number = _equation_number(label)
        if display:
            self._ensure_item_started()
            formula, tag_number = _split_tag(formula)
            number = number or tag_number
        if not formula:
            if display:
                # Keep the empty block so an equation label still has a home.
                paragraph = self._add_paragraph()
                self._decorate(paragraph, "equation")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if number:
                    paragraph.add_run(number)
                return
            # An empty inline span means the dollars were literal, as in
            # "x $ $ y"; put them back rather than swallowing them.
            if self._paragraph is None:
                self._paragraph = self._new_paragraph()
            self._append_text(f"{markup}{latex}{markup}", _plain_formatting(), None)
            return
        if not display and self._prose_reason(formula) is not None:
            # This reads as prose, not a formula -- a lone Cyrillic word
            # (shell variables, price ranges) or, more generally, a plain
            # word-adjacent fragment with no mathematical indicator at all
            # ("$PATH and $HOME", "$low to $high"). Keep it verbatim and
            # let _render_math_literal raise the "write \$" warning.
            self._render_math_literal(latex, display=False, markup=markup)
            return
        try:
            math_element = latex_to_omml(formula)
        except UnsupportedLatexError as error:
            self._warn(f'Formula kept as text: "{formula}" ({error})', "formula_unsupported")
            self._render_math_literal(latex, display, markup=markup)
            return
        self._place_math(math_element, display, number=number)

    def _prose_reason(self, latex: str) -> str | None:
        """Why a single-dollar fragment reads as prose, or ``None`` if it doesn't.

        Two independent signals both mean "this is not math, leave it
        alone": a lone Cyrillic word (shell variables, price ranges written
        in Russian text) or, more generally, a fragment with *no*
        mathematical indicator whatsoever -- no backslash command, no
        ``^``/``_``, no digit, no operator -- that still contains a plain,
        space-separated word of two or more letters, e.g. "PATH and" or
        "low to". A single unspaced token such as "n" or "xy" is left
        alone even though it is all letters, since that is exactly how a
        real formula juxtaposes variable names. Text wrapped in
        ``\\text{...}`` (and friends) is stripped first, since that is how
        real LaTeX writes an ordinary word inside a genuine formula.
        """
        stripped = _TEXT_COMMAND.sub("", latex)
        if _CYRILLIC.search(stripped):
            return "contains Cyrillic text"
        if _MATH_INDICATOR.search(stripped):
            return None
        words = stripped.split()
        if len(words) > 1 and any(word.isalpha() and len(word) >= 2 for word in words):
            return "contains no mathematical symbols"
        return None

    def _place_math(self, math_element: Any, display: bool, number: str | None = None) -> None:
        if display:
            paragraph = self._add_paragraph()
            self._decorate(paragraph, "equation")
            if number:
                # Formula centred, number flush right: two tab stops across
                # the available width, the GOST/LaTeX layout of a numbered
                # equation.
                width = int(self._available_width())
                stops = paragraph.paragraph_format.tab_stops
                stops.add_tab_stop(Emu(width // 2), WD_TAB_ALIGNMENT.CENTER)
                stops.add_tab_stop(Emu(width), WD_TAB_ALIGNMENT.RIGHT)
                paragraph.add_run("\t")
                paragraph._p.append(math_element)
                run = paragraph.add_run(f"\t{number}")
                if self._stamp_runs:
                    run.font.name = self.font_name
                    run.font.size = self.font_size
                return
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            math_paragraph = OxmlElement("m:oMathPara")
            properties = OxmlElement("m:oMathParaPr")
            justification = OxmlElement("m:jc")
            justification.set(qn("m:val"), "center")
            properties.append(justification)
            math_paragraph.append(properties)
            math_paragraph.append(math_element)
            paragraph._p.append(math_paragraph)
            return
        if self._paragraph is None:
            self._paragraph = self._new_paragraph()
        if self._link is not None:
            self._link.append(math_element)
        else:
            self._paragraph._p.append(math_element)

    def _render_amsmath(self, token: Any) -> None:
        r"""Render an ``amsmath`` environment as one or more Word equations.

        The plugin delivers the environment complete with its
        ``\begin{...}``/``\end{...}`` wrapper plus ``meta["environment"]``
        (star already stripped). Matrix families go to ``latex_omml``
        untouched because it parses them itself. ``alignat`` and
        ``flalign`` additionally carry a ``{n}`` column-count argument
        right after the opening tag, which is stripped before the body is
        inspected -- every other environment keeps a leading brace group as
        part of its body, so e.g. ``{\bf x}`` flush against
        ``\begin{equation}`` is not mistaken for one.

        The body is tried whole, as a single equation, first. That is what
        lets a construct nested inside it -- a matrix inside an
        ``equation``, say -- survive intact rather than being cut apart by
        the ``\\``/``&`` splitting below, which knows nothing about
        environment boundaries. ``latex_omml`` renders both separators
        itself -- ``\\`` as a stacked equation array, ``&`` as an OMML
        alignment point -- so a multi-line environment normally converts
        whole and stays ONE Word equation: ``gather`` and ``multline`` with
        their lines stacked, ``align`` and friends with their lines stacked
        *and* aligned on the ``&`` column, which is what all of them mean.

        Only if that whole-body conversion fails does the environment fall
        back to splitting on ``\\`` into one centred equation per line,
        with ``&`` stripped. The case that still reaches it is a
        *single-line* ``&`` environment: with no second line to align
        against, ``latex_omml`` reads a lone ``&`` as a probable unescaped
        ampersand and refuses it, so stripping it here is what keeps
        ``\begin{align}a &= b\end{align}`` an equation instead of verbatim
        text. ``equation`` is single-equation by definition and is never
        split this way. If any line of a split fails to convert, the whole
        environment is kept verbatim so nothing of the source is lost to a
        partial rendering.

        ``\tag{..}`` inside the environment becomes the equation number.
        """
        self._ensure_item_started()
        source = token.content
        environment = (token.meta or {}).get("environment", "")
        if environment in _MATRIX_ENVIRONMENTS:
            self._render_math(source, display=True)
            return

        match = _AMSMATH_WRAPPER.match(source.strip())
        body = match.group("body") if match else source
        if match and environment in _COLUMN_ARGUMENT_ENVIRONMENTS:
            body = _COLUMN_ARGUMENT.sub("", body, count=1)

        body, number = _split_tag(body)

        stripped_body = body.strip()
        if not stripped_body:
            return

        try:
            whole_element = latex_to_omml(stripped_body)
        except UnsupportedLatexError as error:
            whole_body_error = error
        else:
            self._place_math(whole_element, display=True, number=number)
            return

        if environment == "equation":
            # ``equation`` is single-equation by definition: if the whole
            # body did not convert, there is nothing left to try splitting
            # on ``\\``.
            self._warn(
                f'Formula kept as text: "{stripped_body}" ({whole_body_error})',
                "formula_unsupported",
            )
            self._render_math_literal(source, display=True)
            return

        lines = [
            _ALIGNMENT_MARKER.sub(" ", line).strip()
            for line in _LINE_BREAK.split(body)
        ]
        lines = [line for line in lines if line]
        if not lines:
            return

        elements = []
        for line in lines:
            try:
                elements.append(latex_to_omml(line))
            except UnsupportedLatexError as error:
                self._warn(f'Formula kept as text: "{line}" ({error})', "formula_unsupported")
                self._render_math_literal(source, display=True)
                return
        for position, element in enumerate(elements):
            self._place_math(
                element, display=True, number=number if position == len(elements) - 1 else None
            )

    def _render_math_literal(self, latex: str, display: bool, markup: str = "$") -> None:
        """Write a formula as verbatim monospace text, preserving every character.

        Inline fallbacks restore the surrounding ``markup`` (``$`` by
        default) so the user sees exactly what they wrote, delimiters
        included, instead of losing them the one time they matter most.
        """
        text = latex.strip("\n")
        if display:
            paragraph = self._add_paragraph()
            self._decorate(paragraph, "equation")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(text)
        else:
            reason = self._prose_reason(latex)
            if reason is not None:
                self._warn(
                    f'Inline math "${latex}$" {reason} and may be '
                    "ordinary prose rather than a formula; write a literal \"$\" "
                    'as "\\$".',
                    "math_prose",
                )
            if self._paragraph is None:
                self._paragraph = self._new_paragraph()
            run = self._paragraph.add_run(f"{markup}{text}{markup}")
        run.font.name = _CODE_FONT
        run.font.size = Pt(max(7.0, round(self._context_font_size().pt * 0.85)))
