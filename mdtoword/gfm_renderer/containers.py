"""Where a paragraph goes and how the containers around it shape it.

Word has no nesting, so "inside a list inside a quote" is expressed purely
as numbering levels, left indents and borders on flat paragraphs. This
module creates paragraphs in the right place (the body, or the native
footnote being filled), decorates them for the open lists, quotes, alerts
and definition descriptions, and answers the size questions that depend on
context (the indent, the width left for content, the font size in effect).
"""

from __future__ import annotations

from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Emu, Length, Pt

from ..ooxml import set_paragraph_borders, set_paragraph_shading
from .constants import (
    _ALERT_BORDER,
    _ALERT_SHADING,
    _DEFINITION_INDENT,
    _QUOTE_BORDER,
    _QUOTE_STEP,
)
from .state import RendererState


class ContainerMixin(RendererState):
    """Create paragraphs and indent/number/border them for their containers."""

    def _add_paragraph(self, style: str | None = None) -> Any:
        """Create a paragraph wherever content currently goes (body or footnote).

        A pending footnote mark (section-mode notes, footnotes nested in
        another footnote) is written at the start of whatever paragraph comes
        first, so a note that opens with code or a formula keeps its mark.
        """
        if self._footnote_target is not None and self._footnotes is not None:
            paragraph = self._footnotes.add_paragraph(self._footnote_target)
            if style is not None and style not in {"List Paragraph", "Quote"}:
                paragraph.style = self._style(style)
        elif style is None:
            paragraph = self.document.add_paragraph()
        else:
            paragraph = self.document.add_paragraph(style=self._style(style))
        if self._pending_footnote_label is not None:
            run = paragraph.add_run(self._pending_footnote_label)
            run.font.superscript = True
            paragraph.add_run(" ")
            self._pending_footnote_label = None
        return paragraph

    def _container_indent(self) -> Length:
        left = 0
        if self._lists:
            left += self._numbering.continuation_indent(self._lists[-1].level)
        left += self._definition_depth * _DEFINITION_INDENT
        left += len(self._quotes) * _QUOTE_STEP
        return Emu(left)

    def _available_width(self) -> Length:
        return Emu(max(int(self._text_width) - int(self._container_indent()), int(Cm(3))))

    def _decorate(self, paragraph: Any, kind: str) -> None:
        """Indent/number/border *paragraph* for the containers it sits in.

        *kind* is ``"item"`` for the first paragraph of a list item (it gets
        the number), ``"body"`` for ordinary text, or the name of a special
        block (``"code"``, ``"equation"``, ``"figure"``, ``"caption"``, …)
        that never takes the GOST first-line indent.
        """
        paragraph_format = paragraph.paragraph_format
        extra = len(self._quotes) * int(_QUOTE_STEP) + self._definition_depth * int(_DEFINITION_INDENT)
        if self._lists:
            state = self._lists[-1]
            text_indent = int(self._numbering.continuation_indent(state.level))
            if kind == "item":
                self._numbering.apply(paragraph, state.num_id, state.level)
                state.has_paragraph = True
                hanging = int(self._numbering.continuation_indent(0)) // 2 or int(Pt(18))
                if extra or self._gost:
                    paragraph_format.left_indent = Emu(text_indent + extra)
                    paragraph_format.first_line_indent = Emu(-min(hanging, int(Pt(18))))
            else:
                paragraph_format.left_indent = Emu(text_indent + extra)
                paragraph_format.first_line_indent = Pt(0)
        elif extra:
            paragraph_format.left_indent = Emu(extra)
            if self._gost:
                paragraph_format.first_line_indent = Pt(0)
        elif kind != "body" and self._gost:
            paragraph_format.first_line_indent = Pt(0)

        if self._quotes:
            if self._quotes[-1] is not None:
                set_paragraph_shading(paragraph, _ALERT_SHADING)
                set_paragraph_borders(paragraph, left=_ALERT_BORDER)
            else:
                set_paragraph_borders(paragraph, left=_QUOTE_BORDER)

    def _ensure_item_started(self) -> None:
        """Give a list item its number even when it does not start with text.

        A list item whose first block is a code block, a formula or a table
        still needs the numbered paragraph Word hangs the number on.
        """
        if self._lists and not self._lists[-1].has_paragraph:
            paragraph = self._add_paragraph("List Paragraph")
            self._decorate(paragraph, "item")

    def _new_paragraph(self) -> Any:
        if self._lists:
            state = self._lists[-1]
            kind = "body" if state.has_paragraph else "item"
            paragraph = self._add_paragraph("List Paragraph")
            self._decorate(paragraph, kind)
        elif self._quotes:
            paragraph = self._add_paragraph("Quote" if self._quotes[-1] is None else None)
            self._decorate(paragraph, "body")
        else:
            paragraph = self._add_paragraph()
            self._decorate(paragraph, "body")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if self._in_term:
            paragraph.paragraph_format.keep_with_next = True
        return paragraph

    def _context_font_size(self) -> Pt:
        if self._heading_level is not None:
            try:
                size = self.document.styles[f"Heading {self._heading_level}"].font.size
            except KeyError:
                size = None
            return Pt(size.pt) if size is not None else self.font_size
        if self._footnote_target is not None:
            try:
                size = self.document.styles["Footnote Text"].font.size
            except KeyError:
                size = None
            return Pt(size.pt) if size is not None else Pt(10)
        if self._cell_context and self._gost:
            return Pt(max(10, self.font_size.pt - 2))
        return self.font_size
