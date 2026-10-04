"""Inline content: text runs, formatting, links and line breaks.

``_render_inline`` walks the children of an ``inline`` token into the
current paragraph (creating one if needed), turning task-list prefixes into
check boxes and tracking emphasis, strike-through, sub/superscript and
highlight as it goes. Images, footnote references, inline math and inline
HTML are handed to the mixins that own them. Table cells are only collected
here; ``tables.py`` renders them once the whole table is known.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any
from urllib.parse import unquote

from docx.enum.text import WD_COLOR_INDEX
from docx.oxml.ns import qn
from docx.shared import Pt

from ..ooxml import add_external_hyperlink, add_internal_hyperlink, set_run_shading
from .constants import _BLACK, _CODE_FONT, _INLINE_CODE_SHADING, _TASK_PREFIX
from .helpers import _inline_plain_text, _plain_formatting
from .state import RendererState


class InlineMixin(RendererState):
    """Render inline tokens as formatted runs, links and breaks."""

    def _render_inline(
        self, children: list[Any], source_path: Path | None, source_content: str
    ) -> None:
        if self._table_row is not None:
            # Table cells are collected and rendered once the whole table is
            # known (column widths depend on every row).
            self._table_cell_children = children
            return
        if self._paragraph is None:
            self._paragraph = self._new_paragraph()

        task_match = _TASK_PREFIX.match(source_content) if self._lists else None
        if task_match:
            self._paragraph.add_run("☒ " if task_match.group(1).lower() == "x" else "☐ ")
            self._skip_task_prefix(children, task_match.group(0))

        formatting = _plain_formatting()
        if self._cell_context and self._table_header:
            formatting["bold"] = True
        self._link = None
        for token in children:
            self._render_inline_token(token, formatting, source_path)
        self._link = None

    def _render_inline_token(
        self, token: Any, formatting: dict[str, bool], source_path: Path | None
    ) -> None:
        link = self._link
        token_type = token.type
        if token_type == "text":
            self._append_text(token.content, formatting, link)
        elif token_type == "softbreak":
            self._line = (self._line or 0) + 1 if self._line else None
            if self.options.line_breaks == "preserve":
                self._append_break(link)
            else:
                self._append_text(" ", formatting, link)
        elif token_type == "hardbreak":
            self._line = (self._line or 0) + 1 if self._line else None
            self._append_break(link)
        elif token_type == "code_inline":
            self._append_text(token.content, {**formatting, "code": True}, link)
        elif token_type in {"em_open", "em_close"}:
            formatting["italic"] = token_type == "em_open"
        elif token_type in {"strong_open", "strong_close"}:
            formatting["bold"] = token_type == "strong_open" or (
                self._cell_context and self._table_header
            )
        elif token_type in {"s_open", "s_close"}:
            formatting["strike"] = token_type == "s_open"
        elif token_type in {"sub_open", "sub_close"}:
            formatting["sub"] = token_type == "sub_open"
        elif token_type in {"sup_open", "sup_close"}:
            formatting["sup"] = token_type == "sup_open"
        elif token_type in {"mark_open", "mark_close"}:
            formatting["mark"] = token_type == "mark_open"
        elif token_type == "link_open":
            self._link = self._open_link(token.attrGet("href") or "", token.attrGet("title"))
        elif token_type == "link_close":
            self._link = None
        elif token_type == "image":
            self._append_image(token, source_path)
        elif token_type == "footnote_ref":
            self._footnote_reference(token)
        elif token_type in {"math_inline", "math_inline_double"}:
            self._render_math(token.content, display=False, markup=token.markup or "$")
        elif token_type == "html_inline":
            self._link = self._html_inline(token.content, formatting, link, source_path)
        elif token.content:
            self._append_text(token.content, formatting, link)

    @staticmethod
    def _skip_task_prefix(children: list[Any], prefix: str) -> None:
        for token in children:
            if token.type == "text":
                token.content = token.content.removeprefix(prefix)
                return

    def _inline_text(self, children: list[Any]) -> str:
        """Flatten inline children to plain text (used for measuring cells)."""
        return _inline_plain_text(children)

    def _open_link(self, href: str, title: str | None) -> Any:
        if self._paragraph is None:
            return None
        if href.startswith("#"):
            anchor = unquote(href[1:])
            name = self._slug_to_bookmark.get(anchor) or self._slug_to_bookmark.get(anchor.lower())
            if name is None:
                self._warn(
                    f"Link target not found in the document: {href} (kept as plain text)",
                    "link_anchor_missing",
                )
                return None
            element = add_internal_hyperlink(self._paragraph, name)
        else:
            element = add_external_hyperlink(self._paragraph, href)
        if title:
            element.set(qn("w:tooltip"), title)
        return element

    def _append_break(self, link: Any) -> None:
        run = self._paragraph.add_run()
        run.add_break()
        if link is not None:
            link.append(run._r)

    def _append_text(
        self, text: str, formatting: dict[str, bool], link_target: Any
    ) -> None:
        if not text:
            return
        run = self._paragraph.add_run(text)
        self._format_run(run, formatting, hyperlink=link_target is not None)
        if link_target is not None and not isinstance(link_target, str):
            link_target.append(run._r)

    def _format_run(self, run: Any, formatting: dict[str, bool], hyperlink: bool = False) -> None:
        # Direct run formatting always beats style formatting in Word, so
        # stamping every run -- including ones inside a heading -- would
        # flatten every heading level back to body size/font/weight
        # regardless of what _configure_styles set on the Heading N style.
        # Only body text (and table cells) is stamped; headings, footnotes
        # and captions keep their style, while an explicit bold/italic
        # request from the Markdown still wins everywhere.
        quiet = (
            self._heading_level is not None
            or self._footnote_target is not None
            or self._caption_context
            or not self._stamp_runs
        )
        if formatting.get("bold") or not quiet:
            run.bold = formatting.get("bold", False)
        if formatting.get("italic") or not quiet:
            run.italic = formatting.get("italic", False)
        if formatting.get("strike") or not quiet:
            run.font.strike = formatting.get("strike", False)
        if formatting.get("sub"):
            run.font.subscript = True
        if formatting.get("sup"):
            run.font.superscript = True
        if formatting.get("mark"):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        if formatting.get("underline") or hyperlink:
            run.underline = True
        if hyperlink:
            run.font.color.rgb = _BLACK
        if formatting.get("code"):
            run.font.name = _CODE_FONT
            run._r.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), _CODE_FONT)
            size = self._context_font_size().pt * 0.9
            run.font.size = Pt(max(7.0, round(size * 2) / 2))
            set_run_shading(run, _INLINE_CODE_SHADING)
        elif not quiet:
            run.font.name = self.font_name
            run.font.size = self._context_font_size()
