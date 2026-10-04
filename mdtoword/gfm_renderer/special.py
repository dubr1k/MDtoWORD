"""Special paragraphs: figures, captions, TOC markers and GitHub alerts.

Some paragraphs mean more than their text. A paragraph holding only an
image becomes a centred figure with a numbered caption; ``Table: ...``
next to a table becomes the table's numbered caption (always placed above
it); a ``[TOC]`` marker becomes the table of contents; an empty paragraph
left behind by an alert marker is dropped. A blockquote opening with
``[!NOTE]`` and friends becomes a shaded alert with a localized title.
"""

from __future__ import annotations

from collections.abc import Sequence
from pathlib import Path
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH

from ..ooxml import add_field, set_keep_lines
from .constants import _ALERT_MARKER, _ALERT_TITLES, _TABLE_CAPTION, _TOC_MARKER
from .state import RendererState


class SpecialParagraphMixin(RendererState):
    """Recognise and render figures, table captions, TOC markers and alerts."""

    def _render_special_paragraph(
        self, tokens: Sequence[Any], index: int, source_path: Path | None
    ) -> int | None:
        if index + 2 >= len(tokens) or tokens[index + 1].type != "inline":
            return None
        inline = tokens[index + 1]
        children = inline.children or []

        if not children and not inline.content.strip():
            # e.g. the "> [!NOTE]" line once its marker became the callout title.
            return index + 3
        if _TOC_MARKER.match(inline.content) and self._footnote_target is None:
            self._insert_toc()
            return index + 3

        meaningful = [
            child for child in children
            if not (child.type == "text" and not child.content.strip())
            and child.type not in {"softbreak", "hardbreak"}
        ]
        if len(meaningful) == 1 and meaningful[0].type == "image":
            self._render_figure(meaningful[0], source_path)
            return index + 3

        if (
            self._footnote_target is None
            and not self._lists
            and not self._quotes
            and _TABLE_CAPTION.match(inline.content)
        ):
            before = index + 3 < len(tokens) and tokens[index + 3].type == "table_open"
            after = index > 0 and tokens[index - 1].type == "table_close" and self._last_table is not None
            if before or after:
                self._render_table_caption(inline, source_path, before_table=before)
                return index + 3
        return None

    def _render_figure(self, image_token: Any, source_path: Path | None) -> None:
        self._ensure_item_started()
        paragraph = self._add_paragraph()
        self._decorate(paragraph, "figure")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._paragraph = paragraph
        embedded = self._append_image(image_token, source_path)
        self._paragraph = None
        caption = (image_token.attrGet("title") or image_token.content or "").strip()
        if not embedded or not caption:
            return
        set_keep_lines(paragraph, keep_together=True, keep_with_next=True)
        self._figure_count += 1
        caption_paragraph = self._add_paragraph("Caption")
        self._decorate(caption_paragraph, "caption")
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label = "Рисунок" if self._russian else "Figure"
        self._caption_prefix(caption_paragraph, label, "Figure", self._figure_count)
        caption_paragraph.add_run(caption)

    def _caption_prefix(self, paragraph: Any, label: str, sequence: str, number: int) -> None:
        paragraph.add_run(f"{label} ")
        add_field(paragraph, f"SEQ {sequence} \\* ARABIC", str(number))
        paragraph.add_run(" — " if self._russian or self._gost else ": ")

    def _render_table_caption(self, inline: Any, source_path: Path | None, before_table: bool) -> None:
        children = list(inline.children or [])
        prefix = _TABLE_CAPTION.match(inline.content)
        if prefix and children and children[0].type == "text":
            children[0].content = children[0].content[len(prefix.group(0)):].lstrip()
        self._table_count += 1
        paragraph = self._add_paragraph("Caption")
        self._decorate(paragraph, "caption")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        set_keep_lines(paragraph, keep_together=True, keep_with_next=True)
        label = "Таблица" if self._russian else "Table"
        self._caption_prefix(paragraph, label, "Table", self._table_count)
        self._paragraph = paragraph
        self._caption_context = True  # captions keep the Caption style's size
        try:
            self._render_inline(children, source_path, "")
        finally:
            self._caption_context = False
        self._paragraph = None
        if not before_table and self._last_table is not None:
            # Captions go above a table (GOST 7.32 and common practice), even
            # when the Markdown puts the caption line after it.
            self._last_table._tbl.addprevious(paragraph._p)
            self._last_table = None

    def _open_blockquote(self, tokens: Sequence[Any], index: int) -> None:
        self._ensure_item_started()
        kind: str | None = None
        if (
            index + 2 < len(tokens)
            and tokens[index + 1].type == "paragraph_open"
            and tokens[index + 2].type == "inline"
        ):
            match = _ALERT_MARKER.match(tokens[index + 2].content)
            if match:
                kind = match.group(1).lower()
                self._strip_alert_marker(tokens[index + 2])
        self._quotes.append(kind)
        if kind is not None:
            english, russian = _ALERT_TITLES[kind]
            title = self._add_paragraph()
            self._decorate(title, "alert_title")
            title.paragraph_format.keep_with_next = True
            run = title.add_run(russian if self._russian else english)
            run.bold = True
            if not tokens[index + 2].children:
                # "> [!NOTE]" followed only by the next paragraph: drop the
                # now-empty first paragraph instead of leaving a blank line.
                tokens[index + 1].hidden = True

    @staticmethod
    def _strip_alert_marker(inline: Any) -> None:
        children = inline.children or []
        if not children or children[0].type != "text":
            return
        children[0].content = _ALERT_MARKER.sub("", children[0].content, count=1)
        if not children[0].content.strip():
            del children[0]
            if children and children[0].type in {"softbreak", "hardbreak"}:
                del children[0]
        inline.content = _ALERT_MARKER.sub("", inline.content, count=1)
