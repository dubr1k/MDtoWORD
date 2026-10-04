"""The title block from front matter and the table of contents.

The title block (title, subtitle, authors, date) opens the document when
front matter provides it. The table of contents is a real ``TOC`` field
whose cached result already lists every heading as an internal link, so it
works before Word ever refreshes it.
"""

from __future__ import annotations

from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from ..ooxml import add_internal_hyperlink, request_field_update_on_open
from .state import RendererState


class TitleTocMixin(RendererState):
    """Render the front-matter title block and the table of contents."""

    def _render_title_block(self) -> None:
        data = self._front_matter

        def text(key: str) -> str:
            value = data.get(key)
            if isinstance(value, list):
                return ", ".join(value)
            return (value or "").strip()

        title = text("title")
        if title:
            paragraph = self._add_paragraph("Title")
            paragraph.add_run(title)
        subtitle = text("subtitle")
        if subtitle:
            paragraph = self._add_paragraph("Subtitle")
            paragraph.add_run(subtitle)
        for key in ("author", "authors", "date"):
            value = text(key)
            if value:
                paragraph = self._add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.add_run(value)

    def _insert_toc(self) -> None:
        """Insert a TOC field, pre-filled with linked headings.

        The cached result lists every heading (levels 1-3) as an internal
        link, so the table of contents is usable straight away -- even in a
        viewer that never updates fields. Word refreshes it (adding page
        numbers) on open because of ``w:updateFields``.
        """
        if self._toc_inserted:
            return
        self._toc_inserted = True
        heading = self._add_paragraph()
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        heading.paragraph_format.first_line_indent = Pt(0)
        heading.paragraph_format.keep_with_next = True
        title_run = heading.add_run(
            ("СОДЕРЖАНИЕ" if self._gost else "Содержание") if self._russian else "Contents"
        )
        title_run.bold = True

        entries = [(level, text, name) for level, text, name in self._headings if level <= 3]
        instruction = ' TOC \\o "1-3" \\h \\z \\u '

        def field_char(kind: str) -> Any:
            run = OxmlElement("w:r")
            char = OxmlElement("w:fldChar")
            char.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                char.set(qn("w:dirty"), "true")
            run.append(char)
            return run

        def instr() -> Any:
            run = OxmlElement("w:r")
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = instruction
            run.append(element)
            return run

        if not entries:
            paragraph = self._add_paragraph()
            paragraph._p.append(field_char("begin"))
            paragraph._p.append(instr())
            paragraph._p.append(field_char("separate"))
            paragraph.add_run(
                "Обновите поле, чтобы построить оглавление" if self._russian
                else "Update this field to build the table of contents"
            )
            paragraph._p.append(field_char("end"))
        else:
            for position, (level, text, name) in enumerate(entries):
                paragraph = self._add_paragraph()
                paragraph.paragraph_format.first_line_indent = Pt(0)
                paragraph.paragraph_format.left_indent = Cm(0.75 * (level - 1))
                paragraph.paragraph_format.space_after = Pt(0)
                if position == 0:
                    paragraph._p.append(field_char("begin"))
                    paragraph._p.append(instr())
                    paragraph._p.append(field_char("separate"))
                link = add_internal_hyperlink(paragraph, name)
                run = paragraph.add_run(text)
                link.append(run._r)
                if position == len(entries) - 1:
                    paragraph._p.append(field_char("end"))
        request_field_update_on_open(self.document)
