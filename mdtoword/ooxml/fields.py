"""Complex fields, the table of contents and refresh-fields-on-open."""

from __future__ import annotations

import re
from typing import TYPE_CHECKING, Any

from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from ._styles import _find_style_id, _styles_element_of
from ._xml import _SETTINGS_SEQUENCE, _replace_ordered

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument


_TOC_LEVELS = re.compile(r"^[1-9]-[1-9]$")
_DEFAULT_TOC_PLACEHOLDER = "Update field to show the table of contents"


def _fld_char_run(kind: str, dirty: bool = False) -> Any:
    run = OxmlElement("w:r")
    attrs = {qn("w:fldCharType"): kind}
    if dirty:
        attrs[qn("w:dirty")] = "true"
    run.append(OxmlElement("w:fldChar", attrs))
    return run


def add_field(
    paragraph: Paragraph, instruction: str, placeholder: str = "", *, dirty: bool = True
) -> Run | None:
    """Append a complex field (begin/instr/separate/result/end) to ``paragraph``.

    ``dirty`` marks the field for refresh when Word opens the document.
    Returns the run holding the placeholder result (``None`` when empty).
    """
    p = paragraph._p
    p.append(_fld_char_run("begin", dirty))
    instr_run = OxmlElement("w:r")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction.strip()} "
    instr_run.append(instr)
    p.append(instr_run)
    p.append(_fld_char_run("separate"))
    result = paragraph.add_run(placeholder) if placeholder else None
    p.append(_fld_char_run("end"))
    return result


def request_field_update_on_open(document: DocxDocument) -> None:
    """Ask Word to refresh all fields (TOC, SEQ, …) when it opens the file."""
    update = OxmlElement("w:updateFields", {qn("w:val"): "true"})
    _replace_ordered(document.settings.element, update, _SETTINGS_SEQUENCE)


def _toc_instruction(levels: str) -> str:
    if not _TOC_LEVELS.match(levels):
        raise ValueError(f"TOC levels must look like '1-3', got {levels!r}")
    return f'TOC \\o "{levels}" \\h \\z \\u'


def _style_toc_title(paragraph: Paragraph) -> None:
    style_id = _find_style_id(_styles_element_of(paragraph.part), "paragraph", "TOC Heading")
    if style_id is not None:
        paragraph._p.get_or_add_pPr().style = style_id
    else:
        for run in paragraph.runs:
            run.bold = True


def insert_toc_before(
    paragraph: Paragraph,
    levels: str = "1-3",
    placeholder: str = _DEFAULT_TOC_PLACEHOLDER,
    *,
    title: str | None = None,
) -> Paragraph:
    """Insert a TOC field paragraph (and optional title) right before ``paragraph``."""
    instruction = _toc_instruction(levels)
    if title:
        _style_toc_title(paragraph.insert_paragraph_before(title))
    toc = paragraph.insert_paragraph_before()
    add_field(toc, instruction, placeholder)
    return toc


def append_toc(
    document: Any,
    levels: str = "1-3",
    placeholder: str = _DEFAULT_TOC_PLACEHOLDER,
    *,
    title: str | None = None,
) -> Paragraph:
    """Append a TOC field paragraph (and optional title) to ``document`` or a container."""
    instruction = _toc_instruction(levels)
    if title:
        _style_toc_title(document.add_paragraph(title))
    toc = document.add_paragraph()
    add_field(toc, instruction, placeholder)
    return toc


def add_toc(
    paragraph_or_container: Any,
    levels: str = "1-3",
    title: str | None = None,
    document: Any = None,
    placeholder: str = _DEFAULT_TOC_PLACEHOLDER,
) -> Paragraph:
    """Insert a TOC before a Paragraph, or append it to a document/container."""
    del document  # accepted for API compatibility; styles come from the part
    if isinstance(paragraph_or_container, Paragraph):
        return insert_toc_before(paragraph_or_container, levels, placeholder, title=title)
    return append_toc(paragraph_or_container, levels, placeholder, title=title)
