"""Native Word footnotes: the footnotes story part and a manager to fill it.

Importing this module registers :class:`FootnotesPart` with python-docx's
``PartFactory`` (and the ``w:footnotes``/``w:footnote`` element classes), so
documents that already carry a footnotes part (a saved result, or a user
template) load it as a story part.
"""

from __future__ import annotations

import copy
from typing import TYPE_CHECKING, Any

from docx.opc.constants import CONTENT_TYPE as CT
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.oxml import serialize_part_xml
from docx.opc.part import Part, PartFactory
from docx.oxml.ns import nsdecls, qn
from docx.oxml.parser import OxmlElement, parse_xml, register_element_cls
from docx.oxml.xmlchemy import BaseOxmlElement
from docx.parts.story import StoryPart
from docx.text.paragraph import Paragraph
from docx.text.run import Run

from ._styles import _default_style_id, _find_style_id, _styles_element_of, _unused_style_id
from ._xml import _SETTINGS_SEQUENCE, _free_partname, _insert_ordered, _is_element, _xml_attr

if TYPE_CHECKING:
    from docx.document import Document as DocxDocument


FOOTNOTE_TEXT_STYLE_NAME = "footnote text"
FOOTNOTE_REFERENCE_STYLE_NAME = "footnote reference"
_FOOTNOTE_TEXT_STYLE_ID = "FootnoteText"
_FOOTNOTE_REFERENCE_STYLE_ID = "FootnoteReference"
# Drawing ids (wp:docPr/@id) in the footnotes story start here so they cannot
# collide with python-docx's per-part numbering of drawings in the body.
_FOOTNOTE_DRAWING_ID_BASE = 1_000_000

_SEPARATOR_PARAGRAPH = (
    '<w:p><w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
    "<w:r><w:{0}/></w:r></w:p>"
)
_DEFAULT_FOOTNOTES_XML = (
    f'<w:footnotes {nsdecls("w", "r", "wp", "a", "pic", "m", "w14")}>'
    '<w:footnote w:type="separator" w:id="-1">'
    + _SEPARATOR_PARAGRAPH.format("separator")
    + "</w:footnote>"
    '<w:footnote w:type="continuationSeparator" w:id="0">'
    + _SEPARATOR_PARAGRAPH.format("continuationSeparator")
    + "</w:footnote></w:footnotes>"
)


def _has_block_content(footnote: Any) -> bool:
    return any(_is_element(child) for child in footnote)


def _footnote_mark_run(reference_style_id: str | None) -> Any:
    run = OxmlElement("w:r")
    if reference_style_id:
        rpr = OxmlElement("w:rPr")
        rpr.append(OxmlElement("w:rStyle", {qn("w:val"): reference_style_id}))
        run.append(rpr)
    run.append(OxmlElement("w:footnoteRef"))
    return run


class CT_Footnotes(BaseOxmlElement):
    """``w:footnotes`` root of the footnotes part."""


class CT_FtnEdn(BaseOxmlElement):
    """``w:footnote``; registered so its ``.xpath()`` understands ``w:``/``r:`` prefixes."""


register_element_cls("w:footnotes", CT_Footnotes)
register_element_cls("w:footnote", CT_FtnEdn)


class FootnotesPart(StoryPart):
    """``/word/footnotes.xml`` as a story part.

    Being a :class:`~docx.parts.story.StoryPart` lets python-docx objects
    living in a footnote use the regular machinery: ``run.add_picture()``
    relates the image to this part and ``paragraph.part.relate_to()`` puts
    hyperlink relationships here.
    """

    @classmethod
    def new(cls, package: Any) -> FootnotesPart:
        partname = _free_partname(package, "/word/footnotes.xml", "/word/footnotes%d.xml")
        return cls(partname, CT.WML_FOOTNOTES, parse_xml(_DEFAULT_FOOTNOTES_XML), package)

    @property
    def blob(self) -> bytes:
        # The schema requires at least one block in every footnote. A reserved
        # footnote that never received content is serialized with just its
        # mark, without touching the in-memory tree.
        element = self._element
        if any(not _has_block_content(fn) for fn in element.iterchildren(qn("w:footnote"))):
            element = copy.deepcopy(element)
            styles = _styles_element_of(self)
            text_style = _find_style_id(styles, "paragraph", FOOTNOTE_TEXT_STYLE_NAME)
            reference_style = _find_style_id(styles, "character", FOOTNOTE_REFERENCE_STYLE_NAME)
            for footnote in element.iterchildren(qn("w:footnote")):
                if not _has_block_content(footnote):
                    paragraph = OxmlElement("w:p")
                    if text_style:
                        paragraph.get_or_add_pPr().style = text_style
                    paragraph.append(_footnote_mark_run(reference_style))
                    footnote.append(paragraph)
        return serialize_part_xml(element)

    @property
    def next_id(self) -> int:
        return max(super().next_id, _FOOTNOTE_DRAWING_ID_BASE)


PartFactory.part_type_for[CT.WML_FOOTNOTES] = FootnotesPart


def _upgrade_to_footnotes_part(document_part: Any, old_part: Part) -> FootnotesPart:
    """Swap a generic ``Part`` (loaded before registration) for a FootnotesPart."""
    new_part = FootnotesPart.load(
        old_part.partname, old_part.content_type, old_part.blob, old_part.package
    )
    new_part.__dict__["rels"] = old_part.rels  # keep its images/hyperlinks
    for rel in document_part.rels.values():
        if not rel.is_external and rel.target_part is old_part:
            rel._target = new_part
    return new_part


class FootnoteManager:
    """Native Word footnotes for a python-docx document.

    The footnotes part, its ``w:footnotePr`` entry in settings.xml and the
    "footnote text"/"footnote reference" styles are created lazily on first
    use, so constructing a manager for a document without footnotes leaves it
    untouched. An existing footnotes part (user template, re-opened result) is
    reused and new ids continue after the highest one present.
    """

    def __init__(self, document: DocxDocument) -> None:
        self._document = document
        self._part: FootnotesPart | None = None
        self._footnotes: dict[int, Any] = {}
        self.text_style_id: str | None = None
        self.reference_style_id: str | None = None

    @property
    def part(self) -> FootnotesPart:
        """The footnotes part, created (with styles and settings) on first access."""
        if self._part is None:
            self._part = self._ensure_part()
        return self._part

    def reserve(self) -> int:
        """Create an empty footnote and return its id (1, 2, …)."""
        part = self.part
        footnote_id = max([0, *self._footnotes]) + 1
        footnote = OxmlElement("w:footnote", {qn("w:id"): str(footnote_id)})
        part.element.append(footnote)
        self._footnotes[footnote_id] = footnote
        return footnote_id

    def insert_reference(self, paragraph: Paragraph) -> int:
        """Reserve a footnote, append its reference mark to ``paragraph``, return the id."""
        footnote_id = self.reserve()
        self.add_reference(paragraph, footnote_id)
        return footnote_id

    def add_reference(self, paragraph: Paragraph, footnote_id: int) -> Run:
        """Append the superscript reference run for ``footnote_id`` to ``paragraph``."""
        self._footnote(footnote_id)
        run = OxmlElement("w:r")
        if self.reference_style_id:
            rpr = OxmlElement("w:rPr")
            rpr.append(OxmlElement("w:rStyle", {qn("w:val"): self.reference_style_id}))
            run.append(rpr)
        run.append(OxmlElement("w:footnoteReference", {qn("w:id"): str(int(footnote_id))}))
        paragraph._p.append(run)
        return Run(run, paragraph)

    def add_paragraph(
        self, footnote_id: int, style: str | None = "Footnote Text"
    ) -> Paragraph:
        """Append a paragraph to footnote ``footnote_id`` and return it.

        The paragraph's ``.part`` is the footnotes part. The first paragraph
        of a footnote starts with the footnote mark and a space.
        """
        footnote = self._footnote(footnote_id)
        is_first = not _has_block_content(footnote)
        element = OxmlElement("w:p")
        footnote.append(element)
        paragraph = Paragraph(element, self.part)
        if style is not None:
            element.get_or_add_pPr().style = self._resolve_paragraph_style(style)
        if is_first:
            element.append(_footnote_mark_run(self.reference_style_id))
            paragraph.add_run(" ")
        return paragraph

    def has_content(self, footnote_id: int) -> bool:
        """True once footnote ``footnote_id`` holds at least one paragraph/table."""
        return _has_block_content(self._footnote(footnote_id))

    # -- internals ---------------------------------------------------------

    def _footnote(self, footnote_id: int) -> Any:
        self.part  # noqa: B018 -- make sure the part and the index exist
        try:
            return self._footnotes[int(footnote_id)]
        except KeyError:
            raise KeyError(f"no footnote with id {footnote_id}") from None

    def _resolve_paragraph_style(self, style: str) -> str:
        styles = self._document.styles.element
        if style.casefold() in (FOOTNOTE_TEXT_STYLE_NAME, "footnotetext") and self.text_style_id:
            return self.text_style_id
        style_id = _find_style_id(styles, "paragraph", style)
        if style_id is None:
            raise KeyError(f"no paragraph style named {style!r}")
        return style_id

    def _ensure_part(self) -> FootnotesPart:
        document_part = self._document.part
        try:
            part = document_part.part_related_by(RT.FOOTNOTES)
        except KeyError:
            part = FootnotesPart.new(document_part.package)
            document_part.relate_to(part, RT.FOOTNOTES)
        if not isinstance(part, FootnotesPart):
            part = _upgrade_to_footnotes_part(document_part, part)
        self._ensure_settings()
        self._ensure_styles()
        self._footnotes = {}
        for footnote in part.element.iterchildren(qn("w:footnote")):
            raw_id = footnote.get(qn("w:id")) or ""
            if raw_id.lstrip("-").isdigit():
                self._footnotes[int(raw_id)] = footnote
        return part

    def _ensure_settings(self) -> None:
        settings = self._document.settings.element
        if settings.find(qn("w:footnotePr")) is not None:
            return
        footnote_pr = OxmlElement("w:footnotePr")
        for separator_id in ("-1", "0"):
            footnote_pr.append(OxmlElement("w:footnote", {qn("w:id"): separator_id}))
        _insert_ordered(settings, footnote_pr, _SETTINGS_SEQUENCE)

    def _ensure_styles(self) -> None:
        styles = self._document.styles.element
        text_id = _find_style_id(styles, "paragraph", FOOTNOTE_TEXT_STYLE_NAME)
        if text_id is None:
            text_id = _unused_style_id(styles, _FOOTNOTE_TEXT_STYLE_ID)
            based_on = _default_style_id(styles, "paragraph")
            styles.append(parse_xml(
                f'<w:style {nsdecls("w")} w:type="paragraph" w:styleId="{text_id}">'
                f'<w:name w:val="{FOOTNOTE_TEXT_STYLE_NAME}"/>'
                + (f'<w:basedOn w:val="{_xml_attr(based_on)}"/>' if based_on else "")
                + '<w:uiPriority w:val="99"/><w:semiHidden/><w:unhideWhenUsed/>'
                '<w:pPr><w:spacing w:before="0" w:after="0" w:line="240" w:lineRule="auto"/>'
                '<w:ind w:firstLine="0"/></w:pPr>'
                '<w:rPr><w:sz w:val="20"/><w:szCs w:val="20"/></w:rPr></w:style>'
            ))
        reference_id = _find_style_id(styles, "character", FOOTNOTE_REFERENCE_STYLE_NAME)
        if reference_id is None:
            reference_id = _unused_style_id(styles, _FOOTNOTE_REFERENCE_STYLE_ID)
            based_on = _default_style_id(styles, "character")
            styles.append(parse_xml(
                f'<w:style {nsdecls("w")} w:type="character" w:styleId="{reference_id}">'
                f'<w:name w:val="{FOOTNOTE_REFERENCE_STYLE_NAME}"/>'
                + (f'<w:basedOn w:val="{_xml_attr(based_on)}"/>' if based_on else "")
                + '<w:uiPriority w:val="99"/><w:semiHidden/><w:unhideWhenUsed/>'
                '<w:rPr><w:vertAlign w:val="superscript"/></w:rPr></w:style>'
            ))
        self.text_style_id = text_id
        self.reference_style_id = reference_id
